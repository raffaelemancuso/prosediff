"""Documents of tracked changes (prosediff.tracked): two versions without
tracked changes, compared into a .docx or an .odt whose changes, accepted,
give the new version and, rejected, the old one."""

import re
import zipfile
from pathlib import Path

import docx
import odfdo
import pytest
from docx_plus.revisions import read_revisions
from helpers import NEW, commented_documents, lines, odt_file, pair, word_file

from prosediff import tracked
from prosediff.cli import main
from prosediff.diff import Options, compare_paths
from prosediff.render import check_split, format_of, write_output


def same(a, b, changes):
    """Whether two documents read alike, b's tracked changes settled so."""
    return lines(a, "accept-all") == lines(b, changes)


@pytest.mark.parametrize("fmt", ["docx", "odt"])
def test_accepted_the_new_version_rejected_the_old(tmp_path, fmt):
    """Neither version has tracked changes; the output's changes, accepted,
    read as the new version and, rejected, as the old one."""
    old, new = pair(tmp_path, fmt)
    c = compare_paths(str(old), str(new), Options())
    out = write_output(c, tmp_path / f"out.{fmt}", fmt)
    assert out == tmp_path / f"out.{fmt}"
    assert same(new, out, "accept-all")
    assert same(old, out, "reject-all")


def test_word_revisions_formatting_and_comments(tmp_path):
    """Each change is a revision of its own: words inserted and deleted,
    paragraphs inserted and deleted whole with their marks (no empty
    paragraph left behind), a formatting change; the new comment stays."""
    old, new = pair(tmp_path, "docx", comment=True)
    out = write_output(compare_paths(str(old), str(new), Options()), tmp_path / "o.docx", "docx")
    d = docx.Document(str(out))
    kinds = [(r.revision_type, r.text) for r in read_revisions(d)]
    assert ("insertion", " and aims") in kinds
    assert ("deletion", "small but") in kinds
    assert ("deletion", "Limitations are discussed in the last section.") in kinds
    xml = zipfile.ZipFile(out).read("word/document.xml").decode()
    assert xml.count("<w:rPrChange") == 1  # regional: bold, then italic
    assert xml.count("<w:rPr><w:del ") >= 1
    assert re.search(r'<w:pStyle w:val="ListBullet"/>.*?<w:rPr><w:ins ', xml)
    comments = list(d.comments)
    assert [c.text for c in comments] == ["Say how large."]
    assert comments[0].author == "Anna Rossi"
    styles = [p.style.name for p in d.paragraphs]
    assert styles[0] == "Heading 1" and styles.count("List Bullet") == 3


def test_odt_regions(tmp_path):
    """The .odt holds one region per change, first in the text, a deleted
    paragraph kept in its region with the start of the next."""
    old, new = pair(tmp_path, "odt")
    out = write_output(compare_paths(str(old), str(new), Options()), tmp_path / "o.odt", "odt")
    content = zipfile.ZipFile(out).read("content.xml").decode()
    assert content.index("<text:tracked-changes") < content.index("<text:h")
    assert re.search(
        r"<text:deletion>.*?<text:p[^>]*>Limitations are discussed in the last section\.</text:p>"
        r"<text:list><text:list-item><text:p[^>]*/></text:list-item></text:list></text:deletion>",
        content,
    )


@pytest.mark.parametrize("fmt", ["docx", "odt"])
def test_a_deleted_last_paragraph(tmp_path, fmt):
    """A paragraph deleted after the last one goes with the break before it."""
    make = word_file if fmt == "docx" else odt_file
    old = make(tmp_path / f"old.{fmt}", [("p", "One."), ("p", "Two.")])
    new = make(tmp_path / f"new.{fmt}", [("p", "One.")])
    out = write_output(compare_paths(str(old), str(new), Options()), tmp_path / f"o.{fmt}", fmt)
    assert lines(out, "accept-all") == ["One."]
    assert lines(out, "reject-all") == ["One.", "Two."]


def test_only_a_document_with_one_of_its_kind(tmp_path):
    """Tracked changes are written from a Word document compared with a Word
    document, an OpenDocument text with an OpenDocument text: not from
    Markdown, from a document of the other kind, nor from several files."""
    docs = pair(tmp_path, "docx")
    c = compare_paths(str(docs[0]), str(docs[1]), Options())
    with pytest.raises(ValueError, match="OpenDocument text with an OpenDocument text"):
        write_output(c, tmp_path / "o.odt", "odt")
    odt_new = odt_file(tmp_path / "new.odt", NEW)
    with pytest.raises(ValueError, match=r"the new side is new.odt"):
        write_output(
            compare_paths(str(docs[0]), str(odt_new), Options()), tmp_path / "o.docx", "docx"
        )
    (tmp_path / "a.md").write_text("One.\n", encoding="utf-8")
    (tmp_path / "b.md").write_text("Two.\n", encoding="utf-8")
    md = compare_paths(str(tmp_path / "a.md"), str(tmp_path / "b.md"), Options())
    with pytest.raises(ValueError, match=r"the old side is a.md"):
        write_output(md, tmp_path / "o.docx", "docx")
    for side in ("a", "b"):
        (tmp_path / side).mkdir()
        for name in ("x", "y"):
            word_file(tmp_path / side / f"{name}.docx", [("p", f"{side} {name}.")])
    several = compare_paths(str(tmp_path / "a"), str(tmp_path / "b"), Options())
    with pytest.raises(ValueError, match="2 files changed"):
        write_output(several, tmp_path / "o.docx", "docx")
    tracked.check_paths("a.DOCX", "b.docx", "docx")
    with pytest.raises(ValueError, match=r"the old one is a.md"):
        tracked.check_paths("a.md", "b.docx", "docx")


def test_a_locked_document_is_written_beside_it(tmp_path):
    """Open in Word (locked), the output goes beside it, with a warning."""
    target = tmp_path / "o.docx"
    written = []

    def write(path):
        if path == str(target):
            raise PermissionError(13, "locked")
        written.append(path)

    with pytest.warns(UserWarning, match="locked"):
        other = tracked.save(write, target)
    assert re.fullmatch(r"o__locked_\d{8}_\d{6}\.docx", other.name)
    assert written == [str(other)]


def test_formats_by_name_and_split():
    assert format_of("x.docx") == "docx" and format_of("x.ODT") == "odt"
    check_split("paragraph", "docx")
    with pytest.raises(ValueError, match="paragraph by paragraph"):
        check_split("sentence", "odt")


def test_cli_writes_tracked_changes(tmp_path, capsys):
    old, new = pair(tmp_path, "docx")
    out = tmp_path / "tracked.docx"
    assert main(["--files", str(old), str(new), "-o", str(out)]) == 0
    assert zipfile.is_zipfile(out)
    # two Word documents make no .odt: refused before comparing them
    with pytest.raises(SystemExit):
        main(["--files", str(old), str(new), "--format", "odt"])
    assert "OpenDocument text with an OpenDocument text" in capsys.readouterr().err
    assert not (tmp_path / "old_vs_new.odt").exists()


@pytest.mark.parametrize("fmt", ["docx", "odt"])
def test_the_new_comments_stay_as_written(tmp_path, fmt):
    """The new file's comments are its own, their paragraphs and italics
    as they were."""
    old, new = commented_documents(tmp_path, fmt)
    out = write_output(compare_paths(str(old), str(new), Options()), tmp_path / f"o.{fmt}", fmt)
    if fmt == "docx":
        (comment,) = list(docx.Document(str(out)).comments)
        assert [p.text for p in comment.paragraphs] == [
            "Please cite:",
            "",
            "See Research Policy, 49.",
        ]
        assert [r.text for r in comment.paragraphs[2].runs if r.italic] == ["Research Policy"]
    else:
        content = zipfile.ZipFile(out).read("content.xml").decode()
        note = re.search(r"<office:annotation .*?</office:annotation>", content, re.S)[0]
        assert "Please cite:" in note
        assert re.search(r'<text:span text:style-name="I">Research Policy</text:span>', note)


def rich_word(path, new):
    """A Word document with more than text: landscape, a header, a style of
    its own, an image, a table, and, new, a comment."""
    from docx.enum.section import WD_ORIENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.shared import Cm, RGBColor

    d = docx.Document()
    s = d.sections[0]
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    s.header.paragraphs[0].text = "Working paper"
    d.styles.add_style("Quote Box", WD_STYLE_TYPE.PARAGRAPH).font.color.rgb = RGBColor(0, 0, 0x99)
    d.add_heading("Main results" if new else "Results", level=1)
    d.add_paragraph("The effect is large." if new else "The effect is small.")
    d.add_paragraph("Policy matters.", style="Quote Box")
    logo = Path(tracked.__file__).with_name("logo.png")
    d.add_picture(str(logo), width=Cm(2))
    rows = (
        [("North", "125"), ("South", "80")]
        if new
        else [("North", "120"), ("South", "80"), ("Islands", "15")]
    )
    t = d.add_table(rows=len(rows), cols=2)
    for r, (a, b) in zip(t.rows, rows, strict=True):
        r.cells[0].text, r.cells[1].text = a, b
    last = d.add_paragraph("Data from the patent office.")
    if new:
        d.add_paragraph("A sentence added at the end.")
        d.add_comment(last.runs[0], text="Which office?", author="Anna")
    d.save(path)
    return path


def test_the_new_word_file_is_kept_whole_and_marked(tmp_path):
    """Two Word files are compared into a copy of the new one: its page, header,
    styles, image, table and comment kept; accepting every change gives the
    new one, rejecting every change the old, a row deleted a row again."""
    old = rich_word(tmp_path / "old.docx", False)
    new = rich_word(tmp_path / "new.docx", True)
    out = write_output(compare_paths(str(old), str(new), Options()), tmp_path / "r.docx", "docx")
    d = docx.Document(str(out))
    assert d.sections[0].orientation == docx.enum.section.WD_ORIENT.LANDSCAPE
    assert d.sections[0].header.paragraphs[0].text == "Working paper"
    assert "Quote Box" in [s.name for s in d.styles]
    assert len(d.inline_shapes) == 1 and [c.text for c in d.comments] == ["Which office?"]
    assert len(d.tables[0].rows) == 3  # the deleted row, back as a deleted row
    assert same(new, out, "accept-all") and same(old, out, "reject-all")


def test_the_new_odt_file_is_kept_and_marked(tmp_path):
    """Two OpenDocument texts are compared into a copy of the new one, its
    table and styles kept, its changes marked; a row deleted whole is a row
    of the table again, tracked whole as LibreOffice tracks one."""

    def build(path, new):
        d = odfdo.Document("text")
        body = d.body
        body.clear()
        d.insert_style(odfdo.Style("paragraph", name="Box", area="text", color="#000099"))
        body.append(odfdo.Header(1, "Main results" if new else "Results"))
        body.append(odfdo.Paragraph("The effect is large." if new else "The effect is small."))
        body.append(odfdo.Paragraph("Policy matters.", style="Box"))
        rows = (
            [["North", "125"], ["South", "80"]]
            if new
            else [["North", "120"], ["South", "80"], ["Islands", "15"]]
        )
        table = odfdo.Table("T", width=2, height=len(rows))
        table.set_values(rows)
        body.append(table)
        body.append(odfdo.Paragraph("Data from the patent office."))
        if new:
            body.append(odfdo.Paragraph("A sentence added at the end."))
        d.save(path)
        return path

    old, new = build(tmp_path / "old.odt", False), build(tmp_path / "new.odt", True)
    out = write_output(compare_paths(str(old), str(new), Options()), tmp_path / "r.odt", "odt")
    content = zipfile.ZipFile(out).read("content.xml").decode()
    assert "<table:table " in content and 'text:style-name="Box"' in content
    assert content.count("<table:table-row") == 3  # the deleted row, back
    assert 'loext:text-changes-only="false"' in content
    assert "styles.xml" in zipfile.ZipFile(out).namelist()
    assert "Box" in zipfile.ZipFile(out).read("styles.xml").decode()
    assert same(new, out, "accept-all") and same(old, out, "reject-all")


def test_only_prosediffs_marks_are_left_out_of_deleted_text():
    """Text put back as deleted keeps every character of the document's,
    emoji and ligatures included; only prosediff's marks (the comments'
    placeholders, the footnotes' stand-ins) are left out."""
    from prosediff.comments import END_FIRST, PUA_FIRST
    from prosediff.footnotes import FIRST
    from prosediff.redline import plain_text

    text = "Fine 🙂 ﬁt�"
    assert plain_text(text) == text
    assert plain_text(f"a{chr(PUA_FIRST)}b{chr(END_FIRST)}c{chr(FIRST)}") == "abc"
