"""Comments: folding pandoc comment spans into markers, and the comments panel."""

import pytest
from helpers import END, NOTE, two_folders

from prosediff import Options, compare, compare_paths, render
from prosediff.comments import NEW_COMMENT_MARK, REMOVED_COMMENT_MARK, end_of, number_of
from prosediff.diff import COMMENT_MARK, PLACEHOLDER, Comments, fold_comments, plain, show_comments

MARKER = 'data-author="Anna" data-date="2026-09-23 23:40" data-text="Too long."'


def test_fold_comments_replaces_span_with_one_placeholder():
    """A comment span becomes one placeholder character standing for its
    author, text and date, and the end of its text another, standing for
    the same comment; the same comment with a new id, the same one."""
    comments = Comments()
    folded = fold_comments(f"Text {NOTE}anchor{END} more.", comments)
    assert len(comments) == 1
    assert len(folded) == len("Text XanchorY more.")
    assert folded[12] == end_of(folded[5]) and not PLACEHOLDER.match(folded[12])
    assert number_of(folded[12]) == number_of(folded[5]) == 0
    assert "comment" not in folded
    assert comments.get(folded[5]).author == "Anna"
    assert comments.get(folded[5]).text == "Too long."
    assert comments.get(folded[5]).date == "2026-09-23 23:40"
    renumbered = NOTE.replace('id="3"', 'id="25"')
    renumbered_end = END.replace('id="3"', 'id="25"')
    again = fold_comments(f"Text {renumbered}anchor{renumbered_end} more.", comments)
    assert again == folded and len(comments) == 1


def test_nested_brackets_in_the_note():
    comments = Comments()
    folded = fold_comments('x [see [1] and [2]]{.comment-start id="1" author="A"} y', comments)
    assert folded[2] != "[" and comments.get(folded[2]).text == "see [1] and [2]"


def test_comments_without_text_are_left_out():
    comments = Comments()
    empty = '[]{.comment-start id="4" author="A" date="2026-09-23T10:00:00Z"}'
    blank = '[  ]{.comment-start id="5" author="B"}'
    folded = fold_comments(f"One{empty} two{blank} three{NOTE}.", comments)
    assert len(comments) == 1  # only the comment with text
    assert PLACEHOLDER.sub("", folded) == "One two three."
    # kept on request: a marker each, "(no text)" where the text would be
    comments = Comments()
    folded = fold_comments(f"One{empty} two{blank} three{NOTE}.", comments, keep_empty=True)
    assert len(comments) == 3  # the two empty ones (by A and by B) and the note
    html = str(show_comments(folded, comments))
    assert 'aria-label="comment by A, 2026-09-23 10:00: (no text)"' in html


def test_empty_comments_switch(tmp_path):
    from prosediff.cli import main

    empty = '[]{.comment-start id="4" author="Anna" date="2026-09-23T10:00:00Z"}'
    a, b = two_folders(tmp_path, "p.md", "Text.\n", f"Text.{empty}\n")
    assert compare_paths(a, b).comments == []
    c = compare_paths(a, b, Options(empty_comments=True))
    assert [(e.author, e.text, e.status) for e in c.comments] == [("Anna", "", "new")]
    assert ">(no text)</a>" in render(c)
    # the command line passes the switch on
    out = tmp_path / "r.html"
    assert main(["--folders", str(a), str(b), "--empty-comments", "-o", str(out)]) == 0
    assert ">(no text)</a>" in out.read_text(encoding="utf-8")


def test_other_spans_and_links_untouched():
    text = "[link](http://x) and [word]{.smallcaps} and \\[not\\]{.comment-start}"
    assert fold_comments(text, Comments()) == text


def test_plain_and_show_comments():
    comments = Comments()
    folded = fold_comments(f"a {NOTE}b", comments)
    assert plain(folded) == f"a {COMMENT_MARK}b"
    html = str(show_comments(folded, comments))
    assert f'<span class="comment" tabindex="0" role="note" data-c="0" {MARKER}' in html
    assert 'aria-label="comment by Anna, 2026-09-23 23:40: Too long."' in html
    assert f">{COMMENT_MARK}</span>" in html
    # a comment added since the base has its own icon
    new = str(show_comments(folded, comments, frozenset(PLACEHOLDER.findall(folded))))
    # the balloon, marked new by its class (a green + in the HTML report)
    assert 'class="comment new"' in new and f">{NEW_COMMENT_MARK}</span>" in new
    assert 'aria-label="new comment by Anna' in new
    # and so has a comment removed since the base
    gone = str(show_comments(folded, comments, removed=frozenset(PLACEHOLDER.findall(folded))))
    assert 'class="comment removed"' in gone and f">{REMOVED_COMMENT_MARK}</span>" in gone
    assert 'aria-label="removed comment by Anna' in gone


def test_a_comment_on_both_sides_is_not_shown(tmp_path):
    """A comment whose paragraph was deleted lands on the next one: a
    comment both sides have is left out, moved or not."""
    a, b = two_folders(
        tmp_path,
        "p.md",
        f"First paragraph here.{NOTE}\n\nSecond paragraph here.\n",
        f"{NOTE}Second paragraph here, edited.\n",
    )
    c = compare_paths(a, b)
    (f,) = c.files
    row = next(r for r in f.rows if r.right_no is not None and "Second" in str(r.right))
    assert row.changes == ['added ", edited"']
    assert "comment" not in str(row.right)
    assert not any("comment" in str(r.left) for r in f.rows)
    assert c.comments == []


def test_a_shared_comment_leaves_one_space():
    from prosediff.diff import without_shared_comments

    comments = Comments()
    old = [fold_comments(s, comments) for s in [f"a {NOTE}b", f"c {NOTE}.", f"{NOTE} d", NOTE]]
    new = [fold_comments(f"x{NOTE}", comments)]
    lines, _, labels, _ = without_shared_comments(old, new, ["1", "2", "3", "4"], ["1"])
    assert lines == ["a b", "c.", "d"] and labels == ["1", "2", "3"]


def test_a_paragraph_a_comment_only_moved_into_is_not_shown(tmp_path):
    """The comment was already there, on the paragraph before: its moving
    to the next paragraph is no reason to show either of them."""
    lines = [f"Paragraph {k} of the text." for k in range(40)]
    old = list(lines)
    old[19] += NOTE
    new = list(lines)
    new[20] = NOTE + new[20]
    a, b = two_folders(
        tmp_path,
        "p.md",
        "\n\n".join([*old, "The end."]) + "\n",
        "\n\n".join([*new, "The end, edited."]) + "\n",
    )
    (f,) = compare_paths(a, b).files
    shown = [r for r in f.rows if r.kind != "skip"]
    assert [r.kind for r in shown if r.changed] == ["replace"]
    assert not any("Paragraph 20" in str(r.right) for r in shown)


def test_a_paragraph_with_only_a_new_comment_is_shown(tmp_path):
    """Its text is unchanged, so it is no edit, but it is not folded away
    with the unchanged lines: the new comment shows, marked new, and the
    navigation stops there."""
    lines = [f"Paragraph {k} of the text." for k in range(40)]
    old = "\n\n".join(lines) + "\n"
    lines[20] += NOTE
    c = compare_paths(*two_folders(tmp_path, "p.md", old, "\n\n".join(lines) + "\n"))
    (f,) = c.files
    shown = [r for r in f.rows if r.kind != "skip"]
    row = next(r for r in shown if "Paragraph 20" in str(r.right))
    assert row.kind == "equal" and row.changes == []  # no edit of the text
    assert 'class="comment new"' in str(row.right)
    assert (f.additions, f.deletions) == (0, 0)
    assert [e.status for e in c.comments] == ["new"]
    assert row.first_of_change and f.change_count == 1


def test_compare_fold_comments(builder, tmp_path):
    builder.write_all({"p.md": f"Some text{NOTE} here.{END}\n", "p.txt": f"x {NOTE}\n"})
    base = builder.commit("first")
    renumbered = NOTE.replace('id="3"', 'id="9"')
    builder.write_all({"p.md": f"Some new text{renumbered} here.{END}\n", "p.txt": f"y {NOTE}\n"})
    target = builder.commit("second")
    c = compare(builder.path, base, target)  # folding is the default
    md, txt = c.files
    (row,) = [r for r in md.rows if r.kind == "replace"]
    assert row.changes == ['added "new"']  # the renumbered comment is no change
    assert "comment-start" not in str(row.right)
    assert MARKER not in render(c)  # the comment was already there: not shown
    assert "comment-start" in str(txt.rows[0].right)  # only Markdown is folded
    assert not PLACEHOLDER.search(render(c))
    # without folding the markup is compared as text, and there is no panel
    c = compare(builder.path, base, target, Options(comments="text"))
    html = render(c)
    assert "comment-start" in html
    assert c.comments == [] and '<section class="comments-panel"' not in html


def test_comments_panel_statuses_and_links(tmp_path):
    kept = NOTE
    gone = NOTE.replace("Too long.", "Old remark.")
    new = NOTE.replace("Too long.", "New remark.")
    lines = [f"Line {i}." for i in range(30)]
    old = list(lines)
    old[2] += kept
    old[20] += gone
    cur = list(lines)
    cur[2] += kept
    cur[21] = "Line 21 edited." + new
    a, b = two_folders(tmp_path, "p.md", "\n".join(old) + "\n", "\n".join(cur) + "\n")
    c = compare_paths(a, b)
    status = {e.text: e.status for e in c.comments}
    # the comment both sides have is left out
    assert status == {"Old remark.": "removed", "New remark.": "new"}
    by_text = {e.text: e for e in c.comments}
    assert by_text["New remark."].line == 22
    assert by_text["Old remark."].line == 21
    html = render(c)
    assert "Comments: 1 new, 1 removed</h2>" in html and "Too long." not in html
    for e in c.comments:
        assert f'id="{e.anchor}"' in html
    # the new and the removed comment have their own icons, in the text and
    # in the panel
    assert html.count('class="comment new"') == 1
    assert html.count('class="comment removed"') == 1
    assert by_text["New remark."].icon == NEW_COMMENT_MARK
    assert by_text["Old remark."].icon == REMOVED_COMMENT_MARK
    assert not PLACEHOLDER.search(html)  # every comment became a marker


def test_no_placeholder_left_where_moved_passages_are_hidden(tmp_path):
    """A row with a moved passage keeps how it looked without it, shown when
    moved passages are hidden: its comments become markers too."""
    sentence = "The committee met twice in March to review the draft budget."
    rest = "\n\nkeep\n\nNumbers came."
    old = f"Opening remarks were brief. {sentence} Then everyone left.{rest}\n"
    new = f"Opening remarks{NOTE} were brief. Then everyone left.{rest} {sentence}\n"
    c = compare_paths(*two_folders(tmp_path, "p.md", old, new))
    (f,) = c.files
    assert any(r.without_passages for r in f.rows)
    html = render(c)
    assert 'class="pv-off"' in html
    assert not PLACEHOLDER.search(html)


def test_a_comment_dated_without_a_time_keeps_its_date():
    """A date without a time is shown as it is, like a tracked change's."""
    comments = Comments()
    span = '[Why?]{.comment-start id="1" author="A" date="2026-09-23"}'
    folded = fold_comments(f"x {span} y", comments)
    assert comments.get(folded[2]).date == "2026-09-23"


def commented_documents(tmp_path, ext):
    """Two versions of a document, the new one with a comment of two
    paragraphs, a blank one between them, a title in italics."""
    import docx
    import odfdo

    if ext == "docx":
        for name, comment in (("old", False), ("new", True)):
            d = docx.Document()
            p = d.add_paragraph("Some text.")
            if comment:
                c = d.add_comment(p.runs[0], text="Please cite:", author="Anna")
                c.add_paragraph("")
                second = c.add_paragraph("See ")
                second.add_run("Research Policy").italic = True
                second.add_run(", 49.")
            d.save(tmp_path / f"{name}.docx")
    else:
        for name, comment in (("old", False), ("new", True)):
            d = odfdo.Document("text")
            d.body.clear()
            d.insert_style(odfdo.Style("text", name="I", italic=True), automatic=True)
            p = odfdo.Paragraph("Some text.")
            if comment:
                note = odfdo.Annotation("Please cite:", creator="Anna", name="n1")
                note.append(odfdo.Paragraph(""))
                second = odfdo.Paragraph("See ")
                second.append(odfdo.Span("Research Policy", style="I"))
                second.append(", 49.")
                note.append(second)
                p.insert(note, position=0)
            d.body.append(p)
            d.save(tmp_path / f"{name}.{ext}")
    return tmp_path / f"old.{ext}", tmp_path / f"new.{ext}"


@pytest.mark.parametrize("ext", ["docx", "odt"])
def test_a_comment_keeps_its_paragraphs_and_formatting(tmp_path, ext):
    """A comment of a Word or OpenDocument document keeps its paragraphs,
    the blank line between them and its italics, for the tooltip
    (data-rich) and the panel; its text stays one line, for the other
    formats and to be told apart."""
    old, new = commented_documents(tmp_path, ext)
    c = compare_paths(old, new, Options())
    (entry,) = c.comments
    assert entry.text == "Please cite: See Research Policy, 49."
    assert entry.rich == (
        (("Please cite:", ()),),
        (),
        (("See ", ()), ("Research Policy", ("em",)), (", 49.", ())),
    )
    assert str(entry.html) == "Please cite:<br><br>See <i>Research Policy</i>, 49."
    html = render(c)
    rich = "[[[&#34;Please cite:&#34;, &#34;&#34;]], [], [[&#34;See &#34;, &#34;&#34;]"
    assert f'data-rich="{rich}' in html
    assert "<i>Research Policy</i>" in html  # in the panel
