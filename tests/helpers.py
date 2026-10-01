"""Helpers shared by several test files."""

import re
import zipfile
from io import BytesIO
from xml.sax.saxutils import escape

import docx as python_docx
import odfdo
from lxml import etree

from prosediff import document
from prosediff.assess import Annotation, Assessment
from prosediff.sources import read_document

END = '[]{.comment-end id="3"}'


NOTE = '[Too long.]{.comment-start id="3" author="Anna" date="2026-09-23T23:40:00Z"}'


def docx(path, paragraphs, comment=None):
    """A minimal Word document; paragraphs are lists of (kind, text) runs,
    kind being "run", "ins" or "del"; comment adds one comment on the first
    paragraph."""
    auth = 'w:author="A" w:date="2026-01-01T00:00:00Z"'

    def piece(kind, text):
        t = escape(text)
        if kind == "del":
            return (
                f'<w:del w:id="9" {auth}><w:r><w:delText xml:space="preserve">'
                f"{t}</w:delText></w:r></w:del>"
            )
        return f'<w:ins w:id="8" {auth}>{run(t)}</w:ins>' if kind == "ins" else run(t)

    body = ""
    for k, p in enumerate(paragraphs):
        runs = "".join(piece(kind, text) for kind, text in p)
        if comment and k == 0:
            runs = (
                '<w:commentRangeStart w:id="0"/>' + runs + '<w:commentRangeEnd w:id="0"/>'
                '<w:r><w:commentReference w:id="0"/></w:r>'
            )
        body += f"<w:p>{runs}</w:p>"
    comments = None
    if comment:
        comments = (
            '<w:comment w:id="0" w:author="Anna" w:date="2026-01-01T00:00:00Z">'
            f"<w:p><w:r><w:t>{escape(comment)}</w:t></w:r></w:p></w:comment>"
        )
    return docx_xml(path, body, comments=comments)


def docx_xml(path, body, footnotes=None, comments=None, styles=None):
    """A Word document from raw XML: body is the content of w:body, footnotes
    the w:footnote elements, comments the w:comment elements, styles the
    w:style elements (for what python-docx cannot write: tracked changes,
    footnotes, equations)."""
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    m = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    ns = f'xmlns:w="{w}" xmlns:m="{m}"'
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ct = "application/vnd.openxmlformats-officedocument.wordprocessingml"
    parts = {"document": (f"{ct}.document.main+xml", None)}
    if footnotes is not None:
        parts["footnotes"] = (f"{ct}.footnotes+xml", f"{rel}/footnotes")
    if comments is not None:
        parts["comments"] = (f"{ct}.comments+xml", f"{rel}/comments")
    if styles is not None:
        parts["styles"] = (f"{ct}.styles+xml", f"{rel}/styles")
    types = "".join(
        f'<Override PartName="/word/{name}.xml" ContentType="{t}"/>'
        for name, (t, _) in parts.items()
    )
    doc_rels = "".join(
        f'<Relationship Id="rId{k}" Type="{r}" Target="{name}.xml"/>'
        for k, (name, (_, r)) in enumerate(parts.items())
        if r
    )
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/'
            'vnd.openxmlformats-package.relationships+xml"/>'
            f'<Default Extension="xml" ContentType="application/xml"/>{types}</Types>',
        )
        z.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f'<Relationship Id="rId1" Type="{rel}/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        z.writestr(
            "word/_rels/document.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            f"{doc_rels}</Relationships>",
        )
        z.writestr("word/document.xml", f"<w:document {ns}><w:body>{body}</w:body></w:document>")
        if footnotes is not None:
            z.writestr("word/footnotes.xml", f"<w:footnotes {ns}>{footnotes}</w:footnotes>")
        if comments is not None:
            z.writestr("word/comments.xml", f"<w:comments {ns}>{comments}</w:comments>")
        if styles is not None:
            z.writestr("word/styles.xml", f"<w:styles {ns}>{styles}</w:styles>")
    return path


def kinds(rows):
    return [r.kind for r in rows]


def odt_xml(path, body, styles=""):
    """An OpenDocument text from raw XML: body is the content of office:text,
    styles the automatic styles (what odfdo does not write directly:
    tracked changes, comments with dates)."""
    ns = (
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
        'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
        'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
        'xmlns:xlink="http://www.w3.org/1999/xlink" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" '
        'xmlns:loext="urn:org:documentfoundation:names:experimental:office:xmlns:loext:1.0" '
        'office:version="1.3"'
    )
    mime = "application/vnd.oasis.opendocument.text"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", mime, compress_type=zipfile.ZIP_STORED)
        z.writestr(
            "META-INF/manifest.xml",
            '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:'
            'manifest:1.0" manifest:version="1.3">'
            f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/>'
            '<manifest:file-entry manifest:full-path="content.xml" '
            'manifest:media-type="text/xml"/></manifest:manifest>',
        )
        z.writestr(
            "content.xml",
            f"<office:document-content {ns}><office:automatic-styles>{styles}"
            f"</office:automatic-styles><office:body><office:text>{body}</office:text>"
            "</office:body></office:document-content>",
        )
    return path


def run(text, props=""):
    """A Word run of text, props its w:rPr."""
    return f'<w:r>{props}<w:t xml:space="preserve">{text}</w:t></w:r>'


def strip_tags(html):
    """The text of some HTML, its tags taken out."""
    return re.sub(r"<[^>]+>", "", html)


def two_files(tmp_path, old, new, suffix=".md"):
    """Two versions of a text file, a and b, written as UTF-8."""
    a, b = tmp_path / f"a{suffix}", tmp_path / f"b{suffix}"
    a.write_text(old, encoding="utf-8")
    b.write_text(new, encoding="utf-8")
    return a, b


def two_folders(tmp_path, name, old, new):
    """Two folders, old and new, each holding its version of the file name."""
    for side, text in (("old", old), ("new", new)):
        (tmp_path / side).mkdir()
        (tmp_path / side / name).write_text(text, encoding="utf-8")
    return tmp_path / "old", tmp_path / "new"


def markdown_of(data: bytes, name: str = "a.docx", changes: str = "accept-all") -> str:
    """A Word document's (or, named .odt, an OpenDocument text's) body as
    Markdown, as prosediff reads it."""
    return document.to_markdown(read_document(data, name, changes))


# Two versions of a paper, as Word documents or OpenDocument texts --------------

OLD = [
    ("h", "Introduction"),
    ("p", "Innovation policy has long relied on grants to firms."),
    ("p", "It studies the effect of **regional** subsidies on patenting."),
    ("p", "We find a small but significant effect."),
    ("p", "Limitations are discussed in the last section."),
    ("l", "Patents granted"),
    ("l", "Firms treated"),
]
NEW = [
    ("h", "Introduction and aims"),
    ("p", "Innovation policy has long relied on grants and loans to firms."),
    ("p", "It studies the effect of *regional* subsidies on patenting."),
    ("p", "We find a large and significant effect."),
    ("l", "Patents granted"),
    ("l", "Firms treated, by region"),
    ("l", "Workers hired"),
    ("p", "Robustness checks confirm every result."),
]


def parts(text):
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*|([^*]+)", text):
        yield (m[1] or m[2] or m[3]), bool(m[1]), bool(m[2])


def word_file(path, blocks, comment=False):
    d = python_docx.Document()
    for kind, text in blocks:
        if kind == "h":
            d.add_heading(text, level=1)
            continue
        p = d.add_paragraph(style="List Bullet" if kind == "l" else None)
        for t, bold, italic in parts(text):
            r = p.add_run(t)
            r.bold, r.italic = bold or None, italic or None
    if comment:
        d.add_comment(d.paragraphs[3].runs[0], text="Say how large.", author="Anna Rossi")
    d.save(path)
    return path


def odt_file(path, blocks):
    d = odfdo.Document("text")
    body = d.body
    body.clear()
    d.insert_style(odfdo.Style("text", name="B", bold=True), automatic=True)
    d.insert_style(odfdo.Style("text", name="I", italic=True), automatic=True)
    for kind, text in blocks:
        if kind == "h":
            body.append(odfdo.Header(1, text))
        elif kind == "l":
            last = body.children[-1] if body.children else None
            if last is None or last.tag != "text:list":
                last = odfdo.List()
                body.append(last)
            last.append(odfdo.ListItem(text))
        else:
            p = odfdo.Paragraph()
            for t, bold, italic in parts(text):
                p.append(odfdo.Span(t, style="B" if bold else "I") if bold or italic else t)
            body.append(p)
    d.save(path)
    return path


def pair(tmp_path, ext, comment=False):
    if ext == "docx":
        return (
            word_file(tmp_path / "old.docx", OLD),
            word_file(tmp_path / "new.docx", NEW, comment),
        )
    return odt_file(tmp_path / "old.odt", OLD), odt_file(tmp_path / "new.odt", NEW)


def lines(path, changes):
    """The text prosediff reads from a document, its tracked changes
    accepted or rejected: a line per paragraph."""
    doc = read_document(path.read_bytes(), path.name, changes)
    return [str(line) for line in document.lines(doc, lambda mark: "")]


def commented_documents(tmp_path, ext):
    """Two versions of a document, the new one with a comment of two
    paragraphs, a blank one between them, a title in italics."""
    if ext == "docx":
        for name, comment in (("old", False), ("new", True)):
            d = python_docx.Document()
            p = d.add_paragraph("Some text.")
            if comment:
                c = d.add_comment(p.runs[0], text="Please cite:", author="Anna")
                c.add_paragraph("")
                second = c.add_paragraph("See ")
                second.add_run("Research Policy").italic = True
                second.add_run(", 49.")
            d.save(tmp_path / f"{name}.docx")
    else:
        for name, comment in (("old", False), ("new", True)):
            d = odfdo.Document("text")
            d.body.clear()
            d.insert_style(odfdo.Style("text", name="I", italic=True), automatic=True)
            p = odfdo.Paragraph("Some text.")
            if comment:
                note = odfdo.Annotation("Please cite:", creator="Anna", name="n1")
                note.append(odfdo.Paragraph(""))
                second = odfdo.Paragraph("See ")
                second.append(odfdo.Span("Research Policy", style="I"))
                second.append(", 49.")
                note.append(second)
                p.insert(note, position=0)
            d.body.append(p)
            d.save(tmp_path / f"{name}.{ext}")
    return tmp_path / f"old.{ext}", tmp_path / f"new.{ext}"


# What an AI makes of them, and what that leaves in the documents ---------------

MARKDOWN = (
    "## Verdict\n**Mixed**. The effect grew, unsupported.\n\n## What changed\n- One.\n\n"
    "## Improvements\n- None.\n\n## Problems to fix\n1. The effect."
)
FIXED = Annotation(
    "new", "a large and", "significant effect.", "Nothing supports a large effect.",
    "Say a small effect.", "a small and significant effect.",
)  # fmt: skip
ADVICE = Annotation(
    "new", "Robustness checks confirm", "every result.", "Which checks?", "Name them."
)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
ODF_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
ODF_OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
DC = "http://purl.org/dc/elements/1.1/"


def assessment(notes):
    return Assessment("claude", MARKDOWN, model="opus", annotations=list(notes))


def odt_content(path_or_data):
    """The content.xml of an OpenDocument text (a path, or its bytes)."""
    if isinstance(path_or_data, bytes):
        path_or_data = BytesIO(path_or_data)
    return etree.fromstring(zipfile.ZipFile(path_or_data).read("content.xml"))


def revisions(path):
    body = python_docx.Document(str(path)).element.body
    return [el for el in body.iter() if el.tag in (W + "ins", W + "del")]


def comments_of(path, fmt):
    """Each comment's text, its paragraphs one a line, in the order they
    were made."""
    if fmt == "docx":
        return [
            "\n".join(p.text for p in c.paragraphs)
            for c in python_docx.Document(str(path)).comments
        ]
    out = [
        "\n".join("".join(p.itertext()) for p in a.iter(f"{{{ODF_TEXT}}}p"))
        for a in odt_content(path).iter(f"{{{ODF_OFFICE}}}annotation")
    ]
    return sorted(out, key=lambda t: not t.startswith("AI assessment"))


def authors_of(path, fmt):
    """Who made the document's tracked changes."""
    if fmt == "docx":
        return {r.get(W + "author") for r in revisions(path)}
    return {
        c.text
        for c in odt_content(path).iter(f"{{{DC}}}creator")
        if c.getparent().tag.endswith("change-info")
    }


def co_authored(path, fmt):
    """The new version as a co-author left it: in the passage FIXED fixes,
    "large and " a tracked insertion of theirs and "very " a deletion, and
    a comment of theirs; read with them accepted, it is NEW still."""
    if fmt == "docx":
        d = python_docx.Document(str(path))
        p = next(p for p in d.paragraphs if p.text.startswith("We find"))
        p.runs[0].text = "We find a "
        added, gone, rest = (
            p.add_run("large and "),
            p.add_run("very "),
            p.add_run("significant effect."),
        )
        for run, tag in ((added, "ins"), (gone, "del")):
            mark = python_docx.oxml.OxmlElement(f"w:{tag}")
            mark.set(W + "id", "90" if tag == "ins" else "91")
            mark.set(W + "author", "Anna Rossi")
            mark.set(W + "date", "2026-09-01T10:00:00Z")
            run._r.addprevious(mark)
            mark.append(run._r)
        gone._r.find(W + "t").tag = W + "delText"
        d.add_comment(rest, text="Which effect?", author="Anna Rossi")
        d.save(str(path))
        return path
    z = zipfile.ZipFile(path)
    files = {n: z.read(n) for n in z.namelist()}
    z.close()
    root = etree.fromstring(files["content.xml"])
    t = f"{{{ODF_TEXT}}}"
    p = next(p for p in root.iter(t + "p") if "".join(p.itertext()).startswith("We find"))
    for child in list(p):
        p.remove(child)
    p.text = "We find a "
    start = etree.SubElement(p, t + "change-start", {t + "change-id": "ct1"})
    start.tail = "large and "
    etree.SubElement(p, t + "change-end", {t + "change-id": "ct1"})
    change = etree.SubElement(p, t + "change", {t + "change-id": "ct2"})
    note = etree.SubElement(p, f"{{{ODF_OFFICE}}}annotation")
    etree.SubElement(note, f"{{{DC}}}creator").text = "Anna Rossi"
    etree.SubElement(note, t + "p").text = "Which effect?"
    change.tail = "significant effect."
    body = root.find(f".//{{{ODF_OFFICE}}}text")
    tracked = etree.Element(t + "tracked-changes")
    body.insert(0, tracked)
    for name, kind, words in (("ct1", "insertion", None), ("ct2", "deletion", "very ")):
        region = etree.SubElement(tracked, t + "changed-region", {t + "id": name})
        what = etree.SubElement(region, t + kind)
        info = etree.SubElement(what, f"{{{ODF_OFFICE}}}change-info")
        etree.SubElement(info, f"{{{DC}}}creator").text = "Anna Rossi"
        etree.SubElement(info, f"{{{DC}}}date").text = "2026-09-01T10:00:00"
        if words:
            etree.SubElement(what, t + "p").text = words
    files["content.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    with zipfile.ZipFile(path, "w") as out:
        out.writestr(zipfile.ZipInfo("mimetype"), files.pop("mimetype"))
        for name, data in files.items():
            out.writestr(name, data, zipfile.ZIP_DEFLATED)
    return path


# A backend standing in for the models ------------------------------------------

ANSWER = """## Verdict
**Improves**: the introduction is tighter.

## What changed
- The second sentence was cut.

## Improvements
- Shorter.

## Problems to fix
1. "a claim" is no longer supported.
"""


def fake(answer: str = ANSWER, model: str = "fake-1"):
    """A backend answering answer, recording what it was asked."""
    asked = []

    def runner(backend, system, prompt, model_asked, effort, timeout):
        asked.append((backend, system, prompt, model_asked, effort, timeout))
        return answer, model

    runner.asked = asked
    return runner
