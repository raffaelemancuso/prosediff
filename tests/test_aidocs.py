"""The documents the HTML report offers when the AI marked problems in a
Word document or an OpenDocument text (prosediff.aidocs): the tracked
changes with the AI's comments, and the new version with its fixes."""

import datetime as dt
import re
import zipfile
from io import BytesIO
from pathlib import Path

import docx
import pytest
from test_tracked import lines, pair

from prosediff.aidocs import Place, downloads, fix_of, haystack, locate
from prosediff.assess import Annotation, Assessment
from prosediff.diff import Options, compare_paths
from prosediff.render import render

MARKDOWN = (
    "## Verdict\n**Mixed**. The effect grew, unsupported.\n\n## What changed\n- One.\n\n"
    "## Improvements\n- None.\n\n## Problems to fix\n1. The effect."
)
FIXED = Annotation(
    "new", "a large and", "significant effect.", "Nothing supports a large effect.",
    "Say a small effect.", "a small and significant effect.",
)  # fmt: skip
ADVICE = Annotation(
    "new", "Robustness checks confirm", "every result.", "Which checks?", "Name them."
)
MISSING = Annotation("new", "words nowhere in it", "", "Not there.", "", "none")
OLD_SIDE = Annotation("old", "Limitations are discussed", "", "Removed.", "Keep it.")


def assessment(notes):
    return Assessment("claude", MARKDOWN, model="opus", annotations=list(notes))


def made(tmp_path, fmt, notes=(FIXED, ADVICE, MISSING, OLD_SIDE)):
    old, new = pair(tmp_path, fmt)
    c = compare_paths(str(old), str(new), Options())
    got = downloads(c, assessment(notes))
    for d in got:
        (tmp_path / d.name).write_bytes(d.data)
    return old, new, got


def test_passages_found_as_the_report_finds_them():
    """Spacing, soft hyphens, curly quotes and comments' placeholders do not
    stop a passage being found; the case, only when it must."""
    lines = ["We find a “large”  ef­fect here.", "Next one."]
    note = Annotation("new", 'a "large" effect', "here.", "p")
    assert locate(lines, note) == Place(0, 8, 0, 33)
    assert locate(lines, Annotation("new", "WE FIND", "", "p")) == Place(0, 0, 0, 7)
    assert locate(lines, Annotation("new", "here. Next", "", "p")) == Place(0, 28, 1, 4)
    assert locate(lines, Annotation("new", "absent", "", "p")) is None
    assert haystack(["a  b", "c"])[0] == "a b c "


def test_a_fix_is_its_words_changed():
    """A fix is the words the replacement changes, whole, from the last;
    none over two paragraphs, or when nothing changes."""
    line = "We find a large and significant effect."
    place = locate([line], FIXED)
    assert fix_of(line, place, FIXED) == [(10, 15, "small")]
    assert fix_of(line, Place(0, 8, 1, 3), FIXED) is None
    number = Annotation("new", "fallen about", "", "p", "", "fallen about 19.6 m")
    line2 = "It has fallen about 39.2 m by then."
    assert fix_of(line2, Place(0, 7, 0, 26), number) == [(20, 24, "19.6")]
    unit = Annotation("new", "9.81", "", "p", "", "9.81 m/s², whatever")
    line3 = "about 9.81 m/s, and more"
    assert fix_of(line3, Place(0, 6, 0, 24), unit) == [(16, 24, "whatever"), (11, 15, "m/s²,")]
    same = Annotation("new", "a large", "", "p", "", "a large")
    assert fix_of(line, locate([line], same), same) is None


@pytest.mark.parametrize("fmt", ["docx", "odt"])
def test_the_two_documents_with_the_ais_comments_and_fixes(tmp_path, fmt):
    """The document of tracked changes still reads as the new version
    accepted and as the old one rejected; it has the verdict and a comment
    for each problem found in the new version, with the change proposed.
    The new version with each fix a tracked change of the AI's: accepted,
    the fixed text; rejected, the new version as it was. A problem without a
    fix is a comment with the change proposed."""
    old, new, got = made(tmp_path, fmt)
    assert [d.name for d in got] == [
        f"new_tracked_with_AI_comments.{fmt}",
        f"new_with_AI_fixes.{fmt}",
    ]
    out = tmp_path / got[0].name
    assert lines(out, "accept-all") == lines(new, "accept-all")
    assert lines(out, "reject-all") == lines(old, "accept-all")
    notes = got[0].notes
    assert sorted(notes) == [-1, 0, 1]  # the verdict, and the two passages found
    assert all(len(n["comments"]) == 1 and not n["changes"] for n in notes.values())
    texts = comments_of(out, fmt)
    assert texts[0].startswith("AI assessment by Claude Code (opus)\nMixed. The effect grew")
    assert "tracked changes" not in texts[0]
    assert "Nothing supports a large effect.\nProposed: Say a small effect." in texts
    assert "Which checks?\nProposed: Name them." in texts

    out = tmp_path / got[1].name
    accepted = lines(out, "accept-all")
    assert "We find a small and significant effect." in accepted
    assert accepted == [s.replace("a large and", "a small and") for s in lines(new, "accept-all")]
    assert lines(out, "reject-all") == lines(new, "accept-all")
    assert got[1].notes[0]["changes"] and not got[1].notes[1]["changes"]
    texts = comments_of(out, fmt)
    assert "fixes it wrote out are tracked changes" in texts[0]
    assert "Nothing supports a large effect." in texts
    assert "Which checks?\nProposed: Name them." in texts
    if fmt == "docx":
        authors = {r.get(W + "author") for r in revisions(out)}
        assert authors == {"Claude Code (opus)"}


def test_overlapping_fixes_the_second_a_comment(tmp_path):
    """Two fixes of one passage: the first made, the second a comment."""
    other = Annotation(
        "new", "large and significant", "", "Too strong.", "Weaken it.", "big and significant"
    )
    _, _, got = made(tmp_path, "docx", (FIXED, other))
    assert got[1].notes[0]["changes"] and not got[1].notes[1]["changes"]
    assert "Too strong.\nProposed: Weaken it." in comments_of(tmp_path / got[1].name, "docx")


def test_none_for_what_is_not_one_document(tmp_path):
    """No document for Markdown, nor without problems marked, nor when the
    assessment failed."""
    a, b = tmp_path / "a.md", tmp_path / "b.md"
    a.write_text("One.\n", encoding="utf-8")
    b.write_text("Two.\n", encoding="utf-8")
    md = compare_paths(str(a), str(b), Options())
    assert downloads(md, assessment([FIXED])) == []
    old, new = pair(tmp_path, "docx")
    c = compare_paths(str(old), str(new), Options())
    assert downloads(c, assessment([])) == []
    assert downloads(c, Assessment("claude", error="no login", annotations=[FIXED])) == []
    assert downloads(c, None) == []


def test_the_report_carries_them(tmp_path):
    """The report holds each document, its name and its problems' ids, and
    the script that saves it; without problems, neither."""
    old, new = pair(tmp_path, "docx")
    c = compare_paths(str(old), str(new), Options())
    html = render(c, assessment=assessment([FIXED]))
    assert html.count('class="ai-document"') == 2
    assert "new_with_AI_fixes.docx" in html
    assert "_e.deflateSync" in html  # fflate, inlined
    plain = render(c, assessment=assessment([]))
    assert 'class="ai-document"' not in plain and "_e.deflateSync" not in plain
    off = render(c, assessment=assessment([FIXED]), documents=False)
    assert 'class="ai-document"' not in off


W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def revisions(path):
    body = docx.Document(str(path)).element.body
    return [el for el in body.iter() if el.tag in (W + "ins", W + "del")]


def comments_of(path, fmt):
    """Each comment's text, its paragraphs one a line, in the order they
    were made."""
    if fmt == "docx":
        return ["\n".join(p.text for p in c.paragraphs) for c in docx.Document(str(path)).comments]
    from lxml import etree

    root = etree.parse(BytesIO(zipfile.ZipFile(path).read("content.xml")))
    office = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
    text = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
    out = []
    for a in root.iter(f"{{{office}}}annotation"):
        out.append("\n".join("".join(p.itertext()) for p in a.iter(f"{{{text}}}p")))
    return sorted(out, key=lambda t: not t.startswith("AI assessment"))


@pytest.mark.parametrize("fmt", ["docx", "odt"])
def test_each_comment_covers_its_passage(tmp_path, fmt):
    """A comment covers its passage's words, no more and no less: in the
    fixed version, as the fix left them."""
    _, _, got = made(tmp_path, fmt, (FIXED,))
    for d, words in zip(
        got, ("a large and significant effect.", "a small and significant effect."), strict=True
    ):
        ids = d.notes[0]["comments"]
        assert len(ids) == 1
        assert covered(d.data, fmt, ids[0]) == words


def covered(data, fmt, cid):
    """The text between a comment's start and its end, what is deleted and
    what the comment says left out."""
    z = zipfile.ZipFile(BytesIO(data))
    if fmt == "docx":
        xml = z.read("word/document.xml").decode()
        start = xml.index(f'<w:commentRangeStart w:id="{cid}"/>')
        part = xml[start : xml.index(f'<w:commentRangeEnd w:id="{cid}"/>')]
        part = re.sub(r"<w:delText[^>]*>.*?</w:delText>", "", part)
    else:
        xml = z.read("content.xml").decode()
        start = xml.index(f'office:name="{cid}"')
        part = xml[
            xml.index("</office:annotation>", start) : xml.index(
                f'<office:annotation-end office:name="{cid}"'
            )
        ]
    return re.sub(r"<[^>]+>", "", part)


@pytest.mark.parametrize(
    ("start", "end", "fixed"),
    [
        ("a large and", "significant", "surely a large and significant"),
        ("We find a", "a large", "We find a large, robust"),
    ],
)
def test_words_put_in_at_a_passages_edge_stay_in_its_comment(tmp_path, start, end, fixed):
    """Words the fix puts in before a passage's first word, or after its
    last, are within its comment, the passage inside its paragraph."""
    note = Annotation("new", start, end, "p", "", fixed)
    _, _, got = made(tmp_path, "docx", (note,))
    d = got[1]
    assert d.notes[0]["changes"]
    assert covered(d.data, "docx", d.notes[0]["comments"][0]) == fixed


# A Word pair with what a paper has besides plain paragraphs, made by pandoc
# from tests/data/ai_rich/v1.md and v2.md (its README says how): an equation
# in a paragraph and one of its own, curly quotes and en dashes, a footnote,
# a list with bold and italic words, a table, a link. v2 brings in errors;
# RICH are the problems a model would mark in it, as it writes them: straight
# quotes and a hyphen where the document has curly ones and an en dash.
RICH_DIR = Path(__file__).parent / "data" / "ai_rich"


def rich_notes() -> list[Annotation]:
    def fix(start, replacement, end=""):
        return Annotation("new", start, end or start, "Wrong.", "Fix it.", replacement)

    return [
        # 0: in an equation: no fix, a comment
        fix("the energy E = mc^(3)", "the energy E = mc^(2)"),
        # 1: after the equation, in its paragraph
        fix("which is huge", "which is enormous"),
        # 2: straight quotes for curly ones
        fix(
            'The so-called "inertial mass" is doubled', 'The so-called "inertial mass" is conserved'
        ),
        # 3: a hyphen for an en dash
        fix("closed system - a result checked", "closed system - a result confirmed"),
        # 4: a list item's bold words
        fix("release the whole of", "release a small part of"),
        # 5: an italic word
        fix("weighs slightly less than", "weighs slightly more than"),
        # 6: a table cell
        fix("watt", "joule"),
        # 7: a link's words
        fix("the first paper", "the original paper"),
        # 8: a footnote
        fix("were abandoned in the", "were refined throughout the"),
        # 9: words put in before a footnote's reference
        fix("by later work.", "by later work, as noted."),
        # 10: the equation of its own
        fix("E^(2) = (pc)^(2)", "E^(2) = (pc)^(3)"),
    ]


# what the rich pair's fixes change, as v2's lines read (RICH, 1 to 9)
RICH_FIXED = [
    ("is huge", "is enormous"),
    ("is doubled", "is conserved"),
    ("result checked", "result confirmed"),
    ("the whole of", "a small part of"),
    ("slightly less", "slightly more"),
    ("| watt", "| joule"),
    ("first paper", "original paper"),
    ("were abandoned in", "were refined throughout"),
    ("later work.", "later work, as noted."),
]


@pytest.mark.parametrize("fmt", ["docx", "odt"])
def test_fixes_in_a_rich_document(tmp_path, fmt):
    """In a document with equations, a footnote, a list, a table and a link,
    each fix changes its words and nothing else, the document's curly quotes
    and dashes kept where the model wrote plain ones; a fix in an equation,
    whose words are not the file's text, is a comment, the change proposed
    in it (an .odt's equations are objects, found nowhere: neither). In a
    Word document the footnote's comment is on the footnote's number."""
    old, new = RICH_DIR / f"v1.{fmt}", RICH_DIR / f"v2.{fmt}"
    c = compare_paths(str(old), str(new), Options())
    got = downloads(c, assessment(rich_notes()))
    out = tmp_path / got[1].name
    out.write_bytes(got[1].data)
    expected = []
    for line in lines(new, "accept-all"):
        for a, b in RICH_FIXED:
            line = line.replace(a, b)
        expected.append(line)
    assert lines(out, "accept-all") == expected
    assert lines(out, "reject-all") == lines(new, "accept-all")
    notes = got[1].notes
    assert all(notes[k]["changes"] for k in range(1, 10))
    texts = comments_of(out, fmt)
    if fmt == "docx":
        assert not notes[0]["changes"] and not notes[10]["changes"]
        assert texts.count("Wrong.\nProposed: Fix it.") == 2  # the equations'
        # the words put in take the formatting of those they replace
        body = docx.Document(str(out)).element.body
        # (the text of each w:t: python-docx's elements repeat it to itertext)
        put = {"".join(t.text for t in el.iter(W + "t")): el for el in body.iter(W + "ins")}
        assert put["a small part"].find(f".//{W}b") is not None
        assert put["more"].find(f".//{W}i") is not None
        # Word takes no comment in a footnote: the footnote's is on its number
        assert "In the footnote: Wrong." in texts
        notes_xml = zipfile.ZipFile(out).read("word/footnotes.xml").decode()
        assert "commentRangeStart" not in notes_xml
        assert texts.count("Wrong.") == 8
    else:
        assert 0 not in notes and 10 not in notes
        assert texts.count("Wrong.") == 9


@pytest.mark.parametrize(
    ("words", "found"),
    [
        ('The so-called "inertial mass"', True),
        ("closed system - a result", True),
        ("closed system -- a result", False),
    ],
)
def test_passages_found_whatever_their_quotes_and_dashes(words, found):
    """A passage quoted with straight quotes or a hyphen is found where the
    text has curly quotes or an en dash, not with two hyphens for one dash."""
    line = "The so-called “inertial mass” is kept in every closed system – a result."
    assert (locate([line], Annotation("new", words, "", "p")) is not None) == found


def test_a_fix_keeps_the_documents_quotes_and_dashes():
    """Only the words changed make a fix: not a quote or a dash written
    plainly."""
    line = "The so-called “inertial mass” is doubled – said."
    note = Annotation(
        "new", "The so-called", "said.", "p", "", 'The so-called "inertial mass" is kept - said.'
    )
    assert fix_of(line, locate([line], note), note) == [(33, 40, "kept")]


ODF_TEXT = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
ODF_OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
DC = "http://purl.org/dc/elements/1.1/"


def co_authored(path, fmt):
    """The new version as a co-author left it: in the passage FIXED fixes,
    "large and " a tracked insertion of theirs and "very " a deletion, and
    a comment of theirs; read with them accepted, it is NEW still."""
    if fmt == "docx":
        d = docx.Document(str(path))
        p = next(p for p in d.paragraphs if p.text.startswith("We find"))
        p.runs[0].text = "We find a "
        added, gone, rest = (
            p.add_run("large and "),
            p.add_run("very "),
            p.add_run("significant effect."),
        )
        for run, tag in ((added, "ins"), (gone, "del")):
            mark = docx.oxml.OxmlElement(f"w:{tag}")
            mark.set(W + "id", "90" if tag == "ins" else "91")
            mark.set(W + "author", "Anna Rossi")
            mark.set(W + "date", "2026-09-01T10:00:00Z")
            run._r.addprevious(mark)
            mark.append(run._r)
        gone._r.find(W + "t").tag = W + "delText"
        d.add_comment(rest, text="Which effect?", author="Anna Rossi")
        d.save(str(path))
        return path
    from lxml import etree

    z = zipfile.ZipFile(path)
    files = {n: z.read(n) for n in z.namelist()}
    z.close()
    root = etree.fromstring(files["content.xml"])
    t = f"{{{ODF_TEXT}}}"
    p = next(p for p in root.iter(t + "p") if "".join(p.itertext()).startswith("We find"))
    for child in list(p):
        p.remove(child)
    p.text = "We find a "
    start = etree.SubElement(p, t + "change-start", {t + "change-id": "ct1"})
    start.tail = "large and "
    etree.SubElement(p, t + "change-end", {t + "change-id": "ct1"})
    change = etree.SubElement(p, t + "change", {t + "change-id": "ct2"})
    note = etree.SubElement(p, f"{{{ODF_OFFICE}}}annotation")
    etree.SubElement(note, f"{{{DC}}}creator").text = "Anna Rossi"
    etree.SubElement(note, t + "p").text = "Which effect?"
    change.tail = "significant effect."
    body = root.find(f".//{{{ODF_OFFICE}}}text")
    tracked = etree.Element(t + "tracked-changes")
    body.insert(0, tracked)
    for name, kind, words in (("ct1", "insertion", None), ("ct2", "deletion", "very ")):
        region = etree.SubElement(tracked, t + "changed-region", {t + "id": name})
        what = etree.SubElement(region, t + kind)
        info = etree.SubElement(what, f"{{{ODF_OFFICE}}}change-info")
        etree.SubElement(info, f"{{{DC}}}creator").text = "Anna Rossi"
        etree.SubElement(info, f"{{{DC}}}date").text = "2026-09-01T10:00:00"
        if words:
            etree.SubElement(what, t + "p").text = words
    files["content.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    with zipfile.ZipFile(path, "w") as out:
        out.writestr(zipfile.ZipInfo("mimetype"), files.pop("mimetype"))
        for name, data in files.items():
            out.writestr(name, data, zipfile.ZIP_DEFLATED)
    return path


def authors_of(path, fmt):
    """Who made the document's tracked changes."""
    if fmt == "docx":
        return {r.get(W + "author") for r in revisions(path)}
    from lxml import etree

    root = etree.parse(BytesIO(zipfile.ZipFile(path).read("content.xml")))
    return {
        c.text for c in root.iter(f"{{{DC}}}creator") if c.getparent().tag.endswith("change-info")
    }


@pytest.mark.parametrize("fmt", ["docx", "odt"])
@pytest.mark.parametrize("alone", [True, False], ids=["review", "compared"])
def test_the_co_authors_changes_and_comments_are_kept(tmp_path, fmt, alone):
    """The file with the AI's fixes is the file itself, the co-authors'
    tracked changes and comments in it as they were, the AI's added: its fix
    of words a co-author put in is a change of its own beside theirs, not
    within it. All accepted, the fixed text."""
    from prosediff.diff import review_file

    old, new = pair(tmp_path, fmt)
    co_authored(new, fmt)
    c = review_file(new, Options()) if alone else compare_paths(str(old), str(new), Options())
    # compared, no document of the changes since the old version: the
    # co-author's are tracked in the new one already
    (got,) = downloads(c, assessment([FIXED, ADVICE]))
    assert got.name == f"new_with_AI_fixes.{fmt}"
    out = tmp_path / got.name
    out.write_bytes(got.data)
    accepted = lines(out, "accept-all")
    assert "We find a small and significant effect." in accepted
    assert accepted == [s.replace("a large and", "a small and") for s in lines(new, "accept-all")]
    # all rejected, the text before the co-author's changes: the words of
    # theirs the AI took out still theirs (in an OpenDocument text, the AI's
    # deletion stacked on their insertion, as LibreOffice writes it)
    assert "We find a very significant effect." in lines(out, "reject-all")
    assert authors_of(out, fmt) == {"Anna Rossi", "Claude Code (opus)"}
    texts = comments_of(out, fmt)
    assert "Which effect?" in texts and "Which checks?\nProposed: Name them." in texts
    if fmt == "docx":  # no insertion within another
        body = docx.Document(str(out)).element.body
        assert not [
            x for x in body.iter(W + "ins") if next(x.iterancestors(W + "ins"), None) is not None
        ]


@pytest.mark.parametrize("fmt", ["docx", "odt"])
def test_the_ais_changes_and_comments_by_the_author_chosen(tmp_path, fmt):
    """With an author chosen, the AI's fixes and comments are by that name,
    not by the AI's."""
    old, new = pair(tmp_path, fmt)
    c = compare_paths(str(old), str(new), Options())
    chosen = assessment([FIXED, ADVICE])
    chosen.author = "Referee 2"
    out = tmp_path / f"fixed.{fmt}"
    out.write_bytes(downloads(c, chosen)[1].data)
    assert authors_of(out, fmt) == {"Referee 2"}
    if fmt == "docx":
        assert {x.author for x in docx.Document(str(out)).comments} == {"Referee 2"}
    else:
        from lxml import etree

        root = etree.parse(BytesIO(zipfile.ZipFile(out).read("content.xml")))
        notes = {
            x.text
            for x in root.iter(f"{{{DC}}}creator")
            if x.getparent().tag.endswith("annotation")
        }
        assert notes == {"Referee 2"}


@pytest.mark.parametrize("fmt", ["docx", "odt"])
def test_the_ais_changes_and_comments_dated_alike_in_local_time(tmp_path, fmt):
    """The AI's tracked changes and comments carry one date, the local time
    as Word and LibreOffice write theirs (Word's w:date with a "Z" all the
    same): not UTC, which they would show shifted by the time zone."""
    _, _, got = made(tmp_path, fmt)
    z = zipfile.ZipFile(BytesIO(got[1].data))
    if fmt == "docx":
        xml = z.read("word/document.xml").decode() + z.read("word/comments.xml").decode()
        dates = set(re.findall(r'w:date="([^"]+)"', xml))
    else:
        dates = set(re.findall(r"<dc:date>([^<]+)</dc:date>", z.read("content.xml").decode()))
    (date,) = dates
    when = dt.datetime.fromisoformat(date.removesuffix("Z"))
    assert abs(dt.datetime.now() - when) < dt.timedelta(minutes=5)
