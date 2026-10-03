"""The documents the HTML report offers to download when an AI marked the
problems of a Word document or an OpenDocument text (prosediff.assess):

- the document of tracked changes (prosediff.redline), each problem the AI
  marked a comment on its passage: what is wrong and the change proposed;
- the new version, each fix the AI wrote out as the passage should read
  (Annotation.replacement) a tracked change of the AI's, what is wrong a
  comment on it; a problem it gave no such fix for, a comment as above.
  It is the file itself, its own tracked changes and comments (the
  co-authors') kept as they are, the AI's added on top; never one change
  within another of the same kind (WordNotes.fix, OdtNotes.fix).

Both have a first comment, on the first paragraph, with the AI's verdict.
Only the new version's passages are marked: a problem in text the old
version alone has is in the report only. A passage is found as the report's
script finds it (whatever the spacing, the soft hyphens, straight or curly
quotes, the comments among its words; else whatever the case), in the
lines as prosediff compared them; one it cannot find has no comment and no
fix, and a fix whose passage runs over more than one paragraph is left a
comment.

Which comments and changes are each problem's is kept with the file
(Download.notes), for the report's script to take out those of the problems
left out before saving it.
"""

import base64
import copy
import difflib
import re
import warnings
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import PurePosixPath

from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.run import Run
from lxml import etree

from prosediff.assess import REACH, TYPOGRAPHIC, Annotation, Assessment
from prosediff.comments import ANY_PLACEHOLDER
from prosediff.diff import Comparison, FileDiff, word_ops
from prosediff.document import spaced
from prosediff.hyphenate import SOFT_HYPHEN
from prosediff.redline import (
    ODF,
    XML_ID,
    OdtRedline,
    WordRedline,
    after_properties,
    atoms,
    by_run,
    located,
    odf,
    odf_atoms,
    odf_drop,
    odf_insert,
    odf_located,
    run_like,
    split_before,
    text_element,
)
from prosediff.sources import suffix_of
from prosediff.tracked import TRACKED_FORMATS, author_of, check_tracked, settled
from prosediff.word import COMMENTS_EXTENDED, W14, W15

# A passage is found whatever typographic variants (TYPOGRAPHIC) it quotes,
# and a fix never changes one into the other.
PLAIN = str.maketrans(TYPOGRAPHIC)
# What holds the words a Word document's own tracked changes put in.
WORD_INSERTIONS = {qn("w:ins"), qn("w:moveTo")}


# Finding a passage ----------------------------------------------------------------


@dataclass(frozen=True)
class Place:
    """Where a passage is in a file's lines: from character o1 of line j1 to
    character o2 (excluded) of line j2."""

    j1: int
    o1: int
    j2: int
    o2: int


def haystack(lines: list[str]) -> tuple[str, list]:
    """The lines' text as passages are matched in it, each character with
    its (line, offset); None for the space put between two lines. Comments'
    placeholders and soft hyphens are left out, curly quotes made straight,
    each run of whitespace one space."""
    text, at = [], []

    def put(c, where):
        c = TYPOGRAPHIC.get(c, c)
        if c.isspace():
            if not text or text[-1] == " ":
                return
            c = " "
        text.append(c)
        at.append(where)

    for j, line in enumerate(lines):
        for o, c in enumerate(line):
            if c == SOFT_HYPHEN or ANY_PLACEHOLDER.match(c):
                continue
            put(c, (j, o))
        put(" ", None)
    return "".join(text), at


def normalised(words: str) -> str:
    return haystack([words])[0].strip()


def locate(lines: list[str], note: Annotation, hay=None) -> Place | None:
    """Where a note's passage is in lines: from its start words to its end
    words (within REACH of them; else the start words alone), matched as
    written, else whatever the case; None when the start words are nowhere."""
    text, at = hay or haystack(lines)
    start, end = normalised(note.start), normalised(note.end or note.start)
    if not start:
        return None
    i = text.find(start)
    if i < 0:
        text, start, end = text.lower(), start.lower(), end.lower()
        i = text.find(start)
        if i < 0:
            return None
    k = text.find(end, i)
    if k < 0 or k - i > REACH:
        k, end = i, start
    first, last = at[i], at[k + len(end) - 1]
    if first is None or last is None:
        return None
    return Place(first[0], first[1], last[0], last[1] + 1)


def verdict_text(assessment: Assessment) -> str:
    """The Verdict section of an assessment, as plain text."""
    text = assessment.verdict_section or ""
    return spaced(re.sub(r"[*_`]", "", text))


def fix_of(line: str, place: Place, note: Annotation) -> list | None:
    """The edits of a note's fix, in line's offsets, from the last: (start,
    end, text put in); None when there is none: no replacement, or a passage
    over more than one line, or one the replacement leaves as it is."""
    if not note.replacement or place.j1 != place.j2:
        return None
    # the passage's characters, without the comments' placeholders
    keep = [k for k in range(place.o1, place.o2) if not ANY_PLACEHOLDER.match(line[k])]
    passage = "".join(line[k] for k in keep)

    def at(i: int) -> int:  # the line offset of passage character i (its end: after the last)
        return keep[i] if i < len(keep) else (keep[-1] + 1 if keep else place.o1)

    def after(i: int) -> int:  # the line offset just after passage character i - 1
        return keep[i - 1] + 1 if i > 0 else place.o1

    edits = []
    for a1, a2, b1, b2 in whole_words(passage, note.replacement):
        start = at(a1)
        end = after(a2) if a2 > a1 else start
        edits.append((start, end, note.replacement[b1:b2]))
    return edits[::-1] or None


def with_fixes(lines: list[str], notes: list) -> tuple[list[str], set[int]]:
    """lines as the notes' fixes leave them, and the indexes of the notes
    fixed: each note of the new side whose passage is found within one line
    and has a replacement (fix_of); a fix overlapping another, already made,
    left out. A document's line keeps its styles, the words put in taking
    those of the character they replace (or follow)."""
    hay = haystack(lines)
    edits: dict[int, list] = {}
    fixed = set()
    for k, note in enumerate(notes):
        if note.side != "new":
            continue
        place = locate(lines, note, hay)
        if place is None:
            continue
        found = fix_of(lines[place.j1], place, note)
        if not found:
            continue
        taken = edits.setdefault(place.j1, [])
        if any((s < e2 and s2 < e) or s == s2 for s, e, _ in found for s2, e2, _ in taken):
            continue
        taken += found
        fixed.add(k)
    out = list(lines)
    for j, line_edits in edits.items():
        line = out[j]
        styles = getattr(line, "styles", None)
        for start, end, text in sorted(line_edits, key=lambda x: (x[0], x[1]), reverse=True):
            if styles is not None:
                at = (
                    styles[start]
                    if start < len(styles)
                    else (styles[-1] if styles else frozenset())
                )
                styles = styles[:start] + [at] * len(text) + styles[end:]
            line = line[:start] + text + line[end:]
        out[j] = out[j].replaced(line, styles) if styles is not None else line
    return out, fixed


def whole_words(a: str, b: str) -> list[tuple[int, int, int, int]]:
    """The changes from a to b (word_ops, curly quotes and dashes taken for
    plain ones), each widened to the whole words it
    touches, those that then meet made one: (a1, a2, b1, b2). A number is
    several of word_ops' words ("39", ".", "2"): 39.2 made 19.6 is then one
    change, not "39" to "19" and "2" to "6"."""

    def word(c: str) -> bool:
        return not c.isspace()

    out: list[list[int]] = []
    # typographic variants alike: a fix keeps the document's quotes and dashes
    for op, a1, a2, b1, b2 in word_ops(a.translate(PLAIN), b.translate(PLAIN)):
        if op == "equal":
            continue
        # a change starting (ending) inside a word takes the rest of it; the
        # words before and after are alike on both sides, so both widen alike
        first = a[a1] if a2 > a1 else b[b1]
        last = a[a2 - 1] if a2 > a1 else b[b2 - 1]
        while a1 > 0 and b1 > 0 and word(first) and word(a[a1 - 1]):
            a1, b1 = a1 - 1, b1 - 1
        while a2 < len(a) and b2 < len(b) and word(last) and word(a[a2]):
            a2, b2 = a2 + 1, b2 + 1
        if out and a1 <= out[-1][1]:
            out[-1][1], out[-1][3] = max(out[-1][1], a2), max(out[-1][3], b2)
        else:
            out.append([a1, a2, b1, b2])
    return [tuple(c) for c in out]


def exact(line: str, text: str) -> list[bool]:
    """For each character of line, whether it is a character of text, one to
    one (aligned): not an equation's, a note reference's or a field's, which
    the line has as text and the file has not as ordinary text."""
    out = [False] * len(line)
    matcher = difflib.SequenceMatcher(None, line, text, autojunk=False)
    for tag, i1, i2, _, _ in matcher.get_opcodes():
        if tag == "equal":
            out[i1:i2] = [True] * (i2 - i1)
    return out


def in_text(line: str, text: str, edits: list) -> bool:
    """Whether each of a fix's edits is where line and the file's text are
    one to one: the words it takes out, and a word beside the place it puts
    words in. An edit elsewhere (in an equation) could go only by guess."""
    ok = exact(line, text)

    def good(i: int) -> bool:
        return 0 <= i < len(line) and ok[i]

    for start, end, _ in edits:
        if end > start:
            if not all(ok[i] for i in range(start, end) if not ANY_PLACEHOLDER.match(line[i])):
                return False
        elif not (good(start - 1) or good(start)):
            return False
    return True


# Word ---------------------------------------------------------------------------


def word_run(paragraphs, line: str, o: int, before: bool = False):
    """The w:r that starts at character o of line (before: that ends just
    before it), split so that one does; None past (before: before) the
    paragraphs' text."""
    _, at, where = located(paragraphs, line)
    x = where[o]
    if x <= 0 if before else x >= len(at):
        return None
    if x < len(at):
        split_before(at[x])
        at = atoms(paragraphs)[1]
    return at[x - before][0]


def note_reference(red: WordRedline, p):
    """The run of the reference to the footnote (endnote) paragraph p is in,
    in the document's body, and its kind ("footnote", "endnote"); None, ""
    for a paragraph of the body, or a note referred to nowhere."""
    for kind in ("footnote", "endnote"):
        note = next(p.iterancestors(qn(f"w:{kind}")), None)
        if note is None:
            continue
        nid = note.get(qn("w:id"))
        for ref in red.new.doc.element.body.iter(qn(f"w:{kind}Reference")):
            if ref.get(qn("w:id")) == nid and ref.getparent().tag == qn("w:r"):
                return ref.getparent(), kind
    return None, ""


def out_of_insertions(el, red: WordRedline) -> None:
    """Take el, an insertion of the AI's (or a comment's reference), out of
    the document's own insertion (or move) it was put in, splitting that in
    two around it: Word nests no insertion in another. (A deletion of the
    AI's stays in it, as Word puts one there.)"""
    while (outer := el.getparent()) is not None and outer.tag in WORD_INSERTIONS:
        before, after = el.getprevious() is not None, el.getnext() is not None
        if before and after:
            rest = copy.copy(outer)  # its attributes, not its content
            rest[:] = []
            rest.set(qn("w:id"), str(red.ids.next()))
            for x in list(el.itersiblings()):
                rest.append(x)
            outer.addnext(rest)
        (outer.addnext if before else outer.addprevious)(el)


def mark_paragraphs(first, last, cid: int) -> None:
    """A comment's range from the start of paragraph first to the end of
    paragraph last, its reference at the end."""
    start = OxmlElement("w:commentRangeStart", attrs={qn("w:id"): str(cid)})
    after_properties(first, start)
    last.append(OxmlElement("w:commentRangeEnd", attrs={qn("w:id"): str(cid)}))
    ref = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rpr.append(OxmlElement("w:rStyle", attrs={qn("w:val"): "CommentReference"}))
    ref.append(rpr)
    ref.append(OxmlElement("w:commentReference", attrs={qn("w:id"): str(cid)}))
    last.append(ref)


class WordNotes:
    """The AI's comments and fixes put in a Word document: red, a
    WordRedline, its changes marked (the document of tracked changes) or
    not (the new version)."""

    def __init__(self, red: WordRedline, author: str) -> None:
        self.red = red
        self.author = author
        self.lines = red.f.new_text

    def comment(self, place: Place, paragraphs: list[str], bold_first: str = "") -> int | None:
        """A comment on a place, its paragraphs of text (the first word of
        the second bold when bold_first names it); its id, None when the
        place is not in the document."""
        red, lines = self.red, self.lines
        # the start first: a split keeps a run's first part in the run, so the
        # end's run, found first, could be left before the start
        starts, ends = red.new_paragraphs[place.j1], red.new_paragraphs[place.j2]
        if not starts or not ends:
            return None
        ref, kind = note_reference(red, starts[0])
        if ref is not None:  # Word takes no comment in a note: on its number
            first = last = ref
            paragraphs = [f"In the {kind}: {paragraphs[0]}", *paragraphs[1:]]
        else:
            first = word_run(starts, lines[place.j1], place.o1)
            last = word_run(ends, lines[place.j2], place.o2, before=True)
        if first is not None and last is not None:
            c = red.new.doc.add_comment(
                [Run(first, None), Run(last, None)], author=self.author, initials="AI"
            )
        else:  # no ordinary text to anchor it to (an equation): its paragraphs
            c = red.new.doc.comments.add_comment(author=self.author, initials="AI")
            mark_paragraphs(starts[0], ends[-1], c.comment_id)
        # python-docx stamps it with the UTC time: the tracked changes' local one
        c._comment_elm.set(qn("w:date"), red.date)
        for k, text in enumerate(paragraphs):
            p = c.paragraphs[0] if k == 0 else c.add_paragraph()
            if k and bold_first and text.startswith(bold_first):
                p.add_run(bold_first).bold = True
                text = text[len(bold_first) :]
            p.add_run(text)
        for ref in red.new.doc.element.body.iter(qn("w:commentReference")):
            if ref.get(qn("w:id")) == str(c.comment_id) and ref.getparent().tag == qn("w:r"):
                out_of_insertions(ref.getparent(), red)
        return c.comment_id

    def fix(self, j: int, edits: list, begin: int) -> list[int]:
        """Make edits (fix_of) in line j's paragraphs as tracked changes of
        the AI's, its passage beginning at character begin; the ids of the
        changes."""
        red, line = self.red, self.lines[j]
        paragraphs = red.new_paragraphs[j]
        _, at, where = located(paragraphs, line)
        first = where[begin]
        ids = []
        # from the last: an edit moves only what comes after it
        for x1, x2, put in [(where[s], where[e], put) for s, e, put in edits]:
            at = atoms(paragraphs)[1]
            for x in (x2, x1):
                if 0 < x < len(at):
                    split_before(at[x])
            at = atoms(paragraphs)[1]
            after = None  # what the words put in go after
            for r, _ in by_run(at[x1:x2]):
                after = red.delete_run(r)
                ids.append(int(after.get(qn("w:id"))))
            if not put:
                continue
            # in the formatting of the words it replaces; else of the word
            # before (the first after, at the start)
            k = x1 if x2 > x1 or x1 == 0 else x1 - 1
            run = run_like(at[k][0] if k < len(at) else None)
            run.append(text_element("w:t", put))
            mark = red.revision("w:ins")
            mark.append(run)
            ids.append(int(mark.get(qn("w:id"))))
            # within the passage's comment, which starts before the run of its
            # first word and ends after the run of its last
            if after is not None:
                after.addnext(mark)
            elif x1 == first and x1 < len(at):
                at[x1][0].addprevious(mark)
            elif x1 > 0:
                at[x1 - 1][0].addnext(mark)
            elif at:
                at[0][0].addprevious(mark)
            else:
                paragraphs[0].append(mark)
            out_of_insertions(mark, red)
        return ids


# OpenDocument ---------------------------------------------------------------------


def odf_place(paragraphs, line: str, o: int) -> int:
    """Where character o of line is in paragraphs' text."""
    return odf_located(paragraphs, line)[2][o]


def odf_put(paragraphs, x: int, elements: list, after: bool = False) -> None:
    """Put elements one after the other before character x of paragraphs'
    text (after the last one, x at its end); after: just after character x
    - 1 instead, in the paragraph that one is in."""
    at = odf_atoms(paragraphs)[1]
    first = elements[0]
    odf_insert(at[:x] if after else at, x, first, paragraphs)
    rest, first.tail = first.tail, None
    last = first
    for el in elements[1:]:
        last.addnext(el)
        last = el
    last.tail = rest


def odf_remove(paragraphs, x1: int, x2: int) -> None:
    """Take characters x1 to x2 (excluded) out of paragraphs' text."""
    at = odf_atoms(paragraphs)[1]
    for node, where, k in reversed(at[x1:x2]):
        if where is None:  # a space element: one space fewer
            count = int(node.get(odf("text:c")) or 1)
            if count > 1:
                node.set(odf("text:c"), str(count - 1))
            else:
                odf_drop(node)
        elif where == "text":
            node.text = (node.text[:k] + node.text[k + 1 :]) or None
        else:
            node.tail = (node.tail[:k] + node.tail[k + 1 :]) or None


def odf_put_around(marks: list, end, start) -> None:
    """Put end just before the first of marks, start just after the last."""
    first, last = marks[0], marks[-1]
    first.addprevious(end)  # the text before the first mark stays before end
    end.tail = None
    last.addnext(start)
    start.tail, last.tail = last.tail, None


class OdtNotes:
    """The AI's comments and fixes put in an OpenDocument text: red, an
    OdtRedline, its changes marked (the document of tracked changes) or not
    (the new version)."""

    def __init__(self, red: OdtRedline, author: str) -> None:
        self.red = red
        self.author = author
        self.lines = red.f.new_text
        self.count = 0
        self.mine: set[str] = set()  # the ids of the AI's changes

    def comment(self, place: Place, paragraphs: list[str], bold_first: str = "") -> str | None:
        """A comment on a place (an office:annotation and its end), its
        paragraphs of text; its name, None when the place is not in the
        document."""
        red, lines = self.red, self.lines
        if not red.new_paragraphs[place.j1] or not red.new_paragraphs[place.j2]:
            return None
        self.count += 1
        name = f"pdai{self.count}"
        note = etree.Element(odf("office:annotation"), nsmap={"dc": ODF["dc"]})
        note.set(odf("office:name"), name)
        etree.SubElement(note, odf("dc:creator")).text = self.author
        etree.SubElement(note, odf("dc:date")).text = red.date.isoformat()
        for text in paragraphs:
            p = etree.SubElement(note, odf("text:p"))
            p.text = text
        end = etree.Element(odf("office:annotation-end"))
        end.set(odf("office:name"), name)
        # the end first: putting the start in first would move it
        last, first = red.new_paragraphs[place.j2], red.new_paragraphs[place.j1]
        odf_put(last, odf_place(last, lines[place.j2], place.o2), [end])
        odf_put(first, odf_place(first, lines[place.j1], place.o1), [note])
        return name

    def fix(self, j: int, edits: list, begin: int) -> list[str]:
        """Make edits (fix_of) in line j's paragraphs as tracked changes of
        the AI's, as LibreOffice lays them out (begin, where the passage
        begins, as for Word); the ids of the changes. Within the document's
        own insertion (a co-author's), the AI's changes split it in two, its
        deletion stacked on it as LibreOffice stacks a deletion of words
        another put in (OASIS OFFICE-4174, in LibreOffice's extended schema):
        rejected, the words come back as the co-author's."""
        red, line = self.red, self.lines[j]
        paragraphs = red.new_paragraphs[j]
        _, _, where = odf_located(paragraphs, line)
        first = where[begin]
        ids = []
        # from the last: an edit moves only what comes after it
        for start, end, put in edits:
            x1, x2 = where[start], where[end]
            under = self.opened_at(paragraphs, x1) if x2 > x1 else []
            marks, deletion = [], None
            if end > start:
                gone = red.words(line, start, end)
                if gone.text_recursive:
                    deletion = red.region("del", [gone])
                    ids.append(deletion)
                    marks.append(red.point("text:change", deletion))
            if put:
                name = red.region("ins")
                ids.append(name)
                span = etree.Element(odf("text:span"))
                span.text = put
                marks += [
                    red.point("text:change-start", name),
                    span,
                    red.point("text:change-end", name),
                ]
            if not marks:
                continue
            self.mine.update(ids)
            # the marks after the last word taken out (or put in before the
            # passage's first word, after the word before elsewhere), in its
            # paragraph, then the words taken out: a table cell left empty
            # has no word of its own to find its place by afterwards
            if x2 > x1 or (x1 > 0 and x1 != first):
                odf_put(paragraphs, x2, marks, after=True)
            else:
                odf_put(paragraphs, x1, marks)
            around = self.opened(marks[0])
            inside = [n for n in under if n in around]
            if deletion is not None and inside:
                self.stack(deletion, inside[-1])
            self.split(marks, around)
            if x2 > x1:
                odf_remove(paragraphs, x1, x2)
        self.drop_empty()
        return ids

    def opened(self, at) -> list[str]:
        """The document's own changes (not the AI's) whose range element at
        is in: between their change-start and their change-end."""
        opened: dict[str, None] = {}  # in order
        for el in self.red.root.iter(odf("text:change-start"), odf("text:change-end"), at.tag):
            if el is at:
                break
            name = el.get(odf("text:change-id"))
            if name in self.mine:
                continue
            if el.tag == odf("text:change-start"):
                opened[name] = None
            elif el.tag == odf("text:change-end"):
                opened.pop(name, None)
        return list(opened)

    def opened_at(self, paragraphs, x: int) -> list[str]:
        """The document's own changes whose range character x of
        paragraphs' text is in."""
        probe = etree.Element(odf("text:bookmark"))
        odf_put(paragraphs, x, [probe])
        names = self.opened(probe)
        odf_drop(probe)
        return names

    def stack(self, deletion: str, name: str) -> None:
        """The AI's deletion on the insertion name: the region of the
        deletion followed by that insertion's, as LibreOffice writes a
        deletion of words another put in."""
        regions = self.red.changed_regions()
        insertion = regions[name].find(odf("text:insertion")) if name in regions else None
        if insertion is not None:
            regions[deletion].append(copy.deepcopy(insertion))

    def drop_empty(self) -> None:
        """The insertions left with no text (a co-author's split around a
        fix, the AI having taken all of its first part out): their marks
        and their regions."""
        regions = self.red.changed_regions()
        for start in list(self.red.root.iter(odf("text:change-start"))):
            end = start.getnext()
            name = start.get(odf("text:change-id"))
            if start.tail or end is None or end.tag != odf("text:change-end"):
                continue
            if end.get(odf("text:change-id")) != name:
                continue
            if name in regions:
                regions[name].getparent().remove(regions[name])
            odf_drop(start)
            odf_drop(end)

    def split(self, marks: list, opened: list[str]) -> None:
        """Split each of the document's own changes opened (whose range the
        marks were put in) in two around them, the second a copy of its
        region: no range of the AI's within another's."""
        red = self.red
        regions = self.red.changed_regions()
        for name in opened:
            region = regions.get(name)
            if region is None:
                continue
            k = 2
            while f"{name}_{k}" in regions or f"{name}_{k}" in red.taken:
                k += 1
            again = copy.deepcopy(region)
            for attribute in (odf("text:id"), XML_ID):
                if again.get(attribute) is not None:
                    again.set(attribute, f"{name}_{k}")
            region.addnext(again)
            red.taken.add(f"{name}_{k}")
            regions[f"{name}_{k}"] = again
            # its end (the one it has yet: after the marks) now the second part's
            for e in red.root.iter(odf("text:change-end")):
                if e.get(odf("text:change-id")) == name:
                    e.set(odf("text:change-id"), f"{name}_{k}")
            odf_put_around(
                marks,
                red.point("text:change-end", name),
                red.point("text:change-start", f"{name}_{k}"),
            )


# The downloads ---------------------------------------------------------------------


@dataclass
class Download:
    """A document offered for download: its file name, what it is, its
    bytes, its format ("docx", "odt"), the ids of each problem's comments
    and changes in it (by the problem's index among the assessment's
    annotations), for the report's script to take out those of the problems
    left out."""

    name: str
    label: str
    data: bytes
    fmt: str
    notes: dict[int, dict[str, list]] = field(default_factory=dict)

    @property
    def base64(self) -> str:
        return base64.b64encode(self.data).decode("ascii")

    @property
    def spec(self) -> dict:
        """What the report's script needs besides the bytes."""
        return {"name": self.name, "fmt": self.fmt, "notes": self.notes}


def document_file(comparison: Comparison) -> tuple[FileDiff, str] | None:
    """The one Word document (or OpenDocument text) of a comparison, compared
    with another of its kind or reviewed alone (Comparison.single), and its
    format; None for anything else."""
    if comparison.single:
        f = next((f for f in comparison.files if f.new_data and not f.binary), None)
        fmt = suffix_of(f.new_path).lstrip(".") if f else ""
        if fmt not in TRACKED_FORMATS or not settled(f, fmt):
            return None
        return f, fmt
    for fmt in TRACKED_FORMATS:
        try:
            f = check_tracked(comparison, fmt)
        except ValueError:
            continue
        return (f, fmt) if settled(f, fmt) else None
    return None


# The end of the comment of a problem whose fix is in the file.
FIX_APPLIED = "Fix already applied: it is the tracked change on this passage."


def comment_text(note: Annotation, fixed: bool) -> list[str]:
    """A problem's comment: what is wrong, the change proposed, and, when
    the fix is in the text as a tracked change, that it is applied."""
    proposed = [f"Proposed: {note.solution}"] if note.solution else []
    return [note.problem, *proposed, *([FIX_APPLIED] if fixed else [])]


def summary(assessment: Assessment, fixes: bool) -> list[str]:
    """The first comment: who assessed the changes, its verdict, and where
    the problems it marked are. No count: the report may leave some out."""
    verdict = verdict_text(assessment) or (assessment.verdict or "").capitalize()
    where = "Each problem it marked is a comment on its passage"
    if fixes:
        where += "; the fixes it wrote out are tracked changes, to accept or reject"
    return [
        f"AI assessment by {assessment.title}",
        *([verdict] if verdict else []),
        f"{where}. Written by an AI: check it against the text.",
    ]


COMMENTS_EXTENDED_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"
)


def resolvable(doc, ids: list) -> None:
    """Each of the comments ids of a Word document given its entry in the
    commentsExtended part ([MS-DOCX] CT_CommentEx), not done: the w14:paraId
    of its last paragraph, for the report to mark it resolved (w15:done) when
    asked. The part is made when the document has none, added to when it
    has one (its co-authors' comments keep theirs)."""
    wanted = {str(i) for i in ids}
    rel = next((r for r in doc.part.rels.values() if r.reltype == COMMENTS_EXTENDED), None)
    if rel is not None:
        part = rel.target_part
        root = part._element if hasattr(part, "_element") else etree.fromstring(part.blob)
    else:
        part = None
        root = etree.Element(f"{W15}commentsEx", nsmap={"w15": W15[1:-1], "w14": W14[1:-1]})
    taken = {e.get(f"{W15}paraId") for e in root.iter(f"{W15}commentEx")}
    serial = 0x1A000000
    for c in doc.comments:
        if str(c.comment_id) not in wanted:
            continue
        paragraphs = c._comment_elm.findall(qn("w:p"))
        if not paragraphs:
            continue
        pid = paragraphs[-1].get(f"{W14}paraId")
        if pid is None:
            while f"{serial:08X}" in taken:
                serial += 1
            pid = f"{serial:08X}"  # below 0x80000000, as MS-DOCX asks
            paragraphs[-1].set(f"{W14}paraId", pid)
        if pid not in taken:
            etree.SubElement(root, f"{W15}commentEx", {f"{W15}paraId": pid, f"{W15}done": "0"})
            taken.add(pid)
    if part is None:
        blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        part = Part(
            PackURI("/word/commentsExtended.xml"), COMMENTS_EXTENDED_TYPE, blob, doc.part.package
        )
        doc.part.relate_to(part, COMMENTS_EXTENDED)
    elif not hasattr(part, "_element"):
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def notes_in(red, notes, author: str, assessment: Assessment, fixes: bool) -> dict:
    """Put the AI's comments (and, with fixes, its fixes) in red's document;
    each problem's ids."""
    word = isinstance(red, WordRedline)
    kind, text_of = (WordNotes, atoms) if word else (OdtNotes, odf_atoms)
    put = kind(red, author)
    lines = put.lines
    hay = haystack(lines)
    places = {k: locate(lines, n, hay) for k, n in enumerate(notes) if n.side == "new"}
    places = {k: p for k, p in places.items() if p is not None}
    edits = {k: fix_of(lines[p.j1], p, notes[k]) for k, p in places.items()} if fixes else {}
    # a fix where the words are not the file's own text (an equation): a comment
    for k, e in edits.items():
        j = places[k].j1
        if e and not in_text(lines[j], text_of(red.new_paragraphs[j])[0], e):
            edits[k] = None
    # no two fixes in one passage: the second a comment only
    taken: list[Place] = []
    for k in sorted(edits, key=lambda k: (places[k].j1, places[k].o1)):
        p = places[k]
        if edits[k] and any(q.j1 == p.j1 and q.o1 < p.o2 and p.o1 < q.o2 for q in taken):
            edits[k] = None
        elif edits[k]:
            taken.append(p)
    ids: dict[int, dict[str, list]] = {}
    first = next(
        (j for j, line in enumerate(lines) if line.strip() and red.new_paragraphs[j]), None
    )
    if first is not None:
        c = put.comment(Place(first, 0, first, len(lines[first])), summary(assessment, fixes))
        if c is not None:
            ids[-1] = {"comments": [c], "changes": []}
    # the comments first: the fixes change the text they are placed by
    for k, p in places.items():
        c = put.comment(p, comment_text(notes[k], bool(edits.get(k))), "Proposed: ")
        ids[k] = {"comments": [c] if c is not None else [], "changes": []}
    # from the end of the document, each fix's lines as they were
    for k in sorted((k for k in edits if edits[k]), key=lambda k: places[k].j1, reverse=True):
        ids[k]["changes"] = put.fix(places[k].j1, edits[k], places[k].o1)
    return ids


def downloads(comparison: Comparison, assessment: Assessment | None) -> list[Download]:
    """The documents the report offers when the AI marked problems in a Word
    document or an OpenDocument text: the tracked changes with the AI's
    comments (when the new version has no tracked changes of its own), and
    the new version with its fixes (a document reviewed alone, that one
    only: it has no changes); none for anything else.
    A document that cannot be made is left out, with a warning: the report
    must not cost it."""
    if assessment is None or assessment.error or not assessment.annotations:
        return []
    found = document_file(comparison)
    if found is None:
        return []
    f, fmt = found
    notes = assessment.annotations
    author = assessment.author.strip() or assessment.title
    stem = PurePosixPath((f.new_path or "document").replace("\\", "/")).stem
    red_of = WordRedline if fmt == "docx" else OdtRedline
    out = []

    def made(what, red_of_it, fixes: bool, name: str, label: str) -> None:
        """The document of red_of_it() with the AI's comments (and its fixes),
        added to out; left out with a warning when it cannot be made: the
        report must not cost it."""
        try:
            red = red_of_it()
            if red is None:
                return
            ids = notes_in(red, notes, author, assessment, fixes)
            if isinstance(red, WordRedline):
                resolvable(red.new.doc, [c for v in ids.values() for c in v["comments"]])
            data = BytesIO()
            red.save(data)
            out.append(Download(name, label, data.getvalue(), fmt, ids))
        except Exception as e:  # a document the redline cannot take
            warnings.warn(f"{what} could not be made: {e}", stacklevel=3)

    def tracked():
        # the changes since the old version: only when the new one has no
        # tracked changes of its own, which would be marked a second time,
        # under another name
        red = red_of(f, author_of(comparison))
        if red.pending:
            return None
        red.run()
        return red

    if not comparison.single:  # a document alone has no changes to track
        made(
            "the tracked changes with the AI's comments",
            tracked,
            False,
            f"{stem}_tracked_with_AI_comments.{fmt}",
            "The changes, tracked, with the AI's comments",
        )
    which = "document" if comparison.single else "new version"
    made(
        f"the {which} with the AI's fixes",
        # the file itself, its own tracked changes (the co-authors') kept as
        # they are, unless the text was read with them rejected or shown
        lambda: red_of(f, author, own=f.document_changes == "accept-all"),
        True,
        f"{stem}_with_AI_fixes.{fmt}",
        f"The {which} with the AI's fixes, tracked",
    )
    return out
