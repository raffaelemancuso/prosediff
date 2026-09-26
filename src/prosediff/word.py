"""Word documents read with python-docx, into paragraphs of styled text.

python-docx opens the document and resolves what the XML alone does not say:
paragraph and run styles, link targets, and the comments (author, date,
text); docx-plus resolves the bold, italic, underline and the like each run
shows, wherever they are set: on the run, in its character style, in the
paragraph's style or the document's defaults. The paragraphs are then walked element
by element, into a prosediff.document.Document, so that everything lands
where it sits in the text:

- headings (Title, Heading 1-6), list paragraphs and tables are blocks of
  their kind (a heading's text is not marked with what its style already
  makes it, a bold heading's text is not also bold), footnotes and endnotes
  numbered blocks of their own, referenced where they are; each run keeps
  its styles, each link its target;
- each comment is kept where it starts, for the rest of prosediff to fold
  into a marker and list in the comments panel;
- tracked changes are settled as asked: accepting keeps the inserted runs
  (w:ins, w:moveTo) and drops the deleted ones (w:del, w:moveFrom),
  rejecting does the reverse, and "all" keeps both, marked as insertions
  and deletions with their author and date. The spaces at the edges of a
  change stay with it, and a comment anchored in dropped text is kept.

No Markdown is involved in the comparison: it is only written, from the
Document, for git's own commands (docx_to_markdown). Headers, footers, text
boxes' layout and page breaks are not text of the body and are left out; an
image is written [image], with its description when it has one, and an
equation as its text.
"""

from collections import Counter
from dataclasses import dataclass, field
from io import BytesIO

import docx
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx_plus.styles.inspect import resolve_effective_formatting

from prosediff.document import (
    EM,
    STRIKE,
    STRONG,
    SUB,
    SUP,
    UNDERLINE,
    Block,
    CommentMark,
    Document,
    Image,
    NoteRef,
    Span,
    Text,
    comments_in,
    comments_only,
    markdown,
    strip,
    to_markdown,
)
from prosediff.language import WordLanguages, most_letters

CHANGES = ("accept", "reject", "all")
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
HEADING_STYLES = {"title": 1, **{f"heading {n}": n for n in range(1, 7)}}
LIST_STYLES = ("list", "bullet", "number")
INSERTED = {qn("w:ins"), qn("w:moveTo")}
DELETED = {qn("w:del"), qn("w:moveFrom")}
# Wrappers whose content is part of the paragraph's text.
TRANSPARENT = {
    qn("w:smartTag"),
    qn("w:customXml"),
    qn("w:fldSimple"),
    qn("w:sdtContent"),
    qn("w:dir"),
    qn("w:bdo"),
}


class WordError(RuntimeError):
    """The file is not a Word document python-docx can read."""


class _Story:
    """What a python-docx paragraph needs of its container: the part it
    belongs to, to resolve styles and link targets."""

    def __init__(self, part) -> None:
        self.part = part


class _NotesPart:
    """The footnotes (or endnotes) part as a paragraph of it sees it: its own
    relationships (the targets of its links), the document's styles, which
    python-docx gives only the document part."""

    def __init__(self, notes_part, document_part) -> None:
        self._notes = notes_part
        self._document = document_part

    @property
    def rels(self):
        return self._notes.rels

    def __getattr__(self, name):
        return getattr(self._document, name)


@dataclass
class Reader:
    document: Document
    changes: str
    comments: dict = field(default_factory=dict)  # id -> python-docx Comment
    notes: dict = field(default_factory=dict)  # ("footnote"/"endnote", id) -> element
    note_order: list = field(default_factory=list)  # [(kind, id)] as referenced
    shown_comments: set = field(default_factory=set)
    # comments of a paragraph deleted as a whole, for the next paragraph
    carried: list = field(default_factory=list)
    # The languages the text is marked with, when asked for, and the letters
    # of the whole in each language.
    languages: WordLanguages | None = None
    letters: Counter = field(default_factory=Counter)
    # The styles of the paragraph being read that its runs do not restate:
    # a heading's own (its text is bold as a heading, not as bold text).
    plain: frozenset = frozenset()

    def language_of(self, paragraphs) -> str | None:
        if self.languages is None:
            return None
        language, counts = self.languages.of(paragraphs)
        self.letters += counts
        return language

    # Paragraph content ----------------------------------------------------------

    def comment(self, cid: str) -> list:
        """Where a comment starts, once."""
        c = self.comments.get(cid)
        if c is None or cid in self.shown_comments:
            return []
        self.shown_comments.add(cid)
        date = c.timestamp.strftime("%Y-%m-%dT%H:%M:%SZ") if c.timestamp else ""
        return [CommentMark(cid, c.author or "", " ".join(c.text.split()), date)]

    def note_ref(self, kind: str, nid: str) -> list:
        key = (kind, nid)
        if key not in self.note_order:
            self.note_order.append(key)
        return [NoteRef(self.note_order.index(key) + 1)]

    def run(self, el, paragraph, deleted: bool, dropping: bool = False) -> list:
        """The inlines of a w:r: its text, in the run's styles.

        dropping: the run is in text that goes (a change settled away); only
        its comments are kept, so its footnote references are not counted.
        """
        r = Run(el, paragraph)
        style = (r.style.name or "").lower() if r.style is not None else ""
        styles = set(styles_of(resolve_effective_formatting(r)))
        if style == "strong":
            styles.add(STRONG)
        elif style == "emphasis":
            styles.add(EM)
        styles = frozenset(styles - self.plain)
        out: list = []
        for child in el:
            tag = child.tag
            if tag == qn("w:t") or (tag == qn("w:delText") and deleted):
                out.append(Text(child.text or "", styles))
            elif tag in (qn("w:tab"), qn("w:ptab"), qn("w:br"), qn("w:cr")):
                out.append(Text(" ", styles))
            elif tag == qn("w:noBreakHyphen"):
                out.append(Text("-", styles))
            elif tag == qn("w:footnoteReference") and not dropping:
                out += self.note_ref("footnote", child.get(qn("w:id")))
            elif tag == qn("w:endnoteReference") and not dropping:
                out += self.note_ref("endnote", child.get(qn("w:id")))
            elif tag == qn("w:commentReference"):
                out += self.comment(child.get(qn("w:id")))
            elif tag in (qn("w:drawing"), qn("w:pict"), qn("w:object")):
                descr = next((d.get("descr") for d in child.iter() if d.get("descr")), "")
                out.append(Image(f"[image: {descr}]" if descr else "[image]"))
        return out

    def children(self, el, paragraph, deleted: bool = False, dropping: bool = False) -> list:
        """The inlines of an element's children, tracked changes settled."""
        out: list = []
        for child in el:
            tag = child.tag
            if tag == qn("w:r"):
                out += self.run(child, paragraph, deleted, dropping)
            elif tag == qn("w:commentRangeStart"):
                out += self.comment(child.get(qn("w:id")))
            elif tag in INSERTED or tag in DELETED:
                is_deletion = tag in DELETED
                keep = self.changes == "all" or (self.changes == "accept") != is_deletion
                inner = self.children(
                    child, paragraph, deleted=is_deletion, dropping=dropping or not keep
                )
                if not keep:
                    # the text goes, the comments anchored in it stay
                    out += comments_in(inner)
                elif self.changes == "all":
                    if markdown(inner):
                        out.append(
                            Span(
                                "deletion" if is_deletion else "insertion",
                                inner,
                                author=child.get(qn("w:author")) or "",
                                date=child.get(qn("w:date")) or "",
                            )
                        )
                else:
                    out += inner
            elif tag == qn("w:hyperlink"):
                inner = self.children(child, paragraph, deleted, dropping)
                rid = child.get(qn("r:id"))
                target = ""
                if rid and rid in paragraph.part.rels:
                    target = paragraph.part.rels[rid].target_ref
                out.append(Span("link", inner, target=target))
            elif tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    out += self.children(content, paragraph, deleted, dropping)
            elif tag in TRANSPARENT:
                out += self.children(child, paragraph, deleted, dropping)
            elif tag in (f"{{{M_NS}}}oMath", f"{{{M_NS}}}oMathPara"):
                out.append(Text(omml_text(child)))
        return out

    # Blocks -----------------------------------------------------------------------

    def paragraph_inlines(self, el, part=None) -> list:
        """A paragraph's inlines, its edges stripped; part is the one it
        belongs to (a footnote's links are looked up among the footnotes'
        relationships)."""
        doc = self.document.part
        p = Paragraph(el, _Story(doc if part is None else _NotesPart(part, doc)))
        style = (p.style.name or "").lower() if p.style is not None else ""
        self.plain = (
            styles_of(resolve_effective_formatting(p)) if style in HEADING_STYLES else frozenset()
        )
        return strip(self.children(el, p))

    def cell(self, paragraphs, part=None) -> list:
        """The inlines of several paragraphs, as one, a space between them."""
        out: list = []
        for p in paragraphs:
            inlines = self.paragraph_inlines(p, part)
            if markdown(inlines):
                if out:
                    out.append(Text(" "))
                out += inlines
        return out

    def paragraph(self, el) -> Block | None:
        inlines = self.paragraph_inlines(el)
        if not markdown(inlines):
            return None
        # A paragraph deleted as a whole keeps only the comments anchored in
        # it: as Word merges it into the next paragraph, they go there.
        if comments_only(inlines):
            self.carried += inlines
            return None
        inlines, self.carried = self.carried + inlines, []
        language = self.language_of([el])
        p = Paragraph(el, _Story(self.document.part))
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style in HEADING_STYLES:
            return Block("heading", inlines, level=HEADING_STYLES[style], language=language)
        numbered = el.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None
        if numbered or style.startswith(LIST_STYLES):
            return Block("item", inlines, language=language)
        return Block("p", inlines, language=language)

    def table(self, el) -> Block:
        rows = [
            [self.cell(tc.iter(qn("w:p"))) for tc in tr.findall(qn("w:tc"))]
            for tr in el.iter(qn("w:tr"))
        ]
        return Block("table", rows=rows, language=self.language_of(el.iter(qn("w:p"))))

    def body(self) -> list[Block]:
        """The paragraphs and tables of the document; comments carried past
        the last paragraph join it."""
        out = self.blocks(self.document.element.body)
        if self.carried:
            if not out:
                out.append(Block("p"))
            last = out[-1]
            if last.kind == "table" and last.rows and last.rows[-1]:
                last.rows[-1][-1] += self.carried
            elif last.kind == "table":
                out.append(Block("p", self.carried))
            else:
                last.inlines += self.carried
            self.carried = []
        return out

    def blocks(self, container) -> list[Block]:
        out: list[Block] = []
        for child in container:
            if child.tag == qn("w:p"):
                if block := self.paragraph(child):
                    out.append(block)
            elif child.tag == qn("w:tbl"):
                out.append(self.table(child))
            elif child.tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    out += self.blocks(content)
            elif child.tag in INSERTED | DELETED or child.tag == qn("w:customXml"):
                out += self.blocks(child)
        return out

    def note_blocks(self) -> list[Block]:
        """The footnotes and endnotes, in the order they are referenced (a note
        may reference another, so the list can grow while it is written)."""
        out = []
        k = 0
        while k < len(self.note_order):
            found = self.notes.get(self.note_order[k])
            k += 1
            if found is None:
                continue
            el, part = found
            paragraphs = list(el.iter(qn("w:p")))
            out.append(
                Block(
                    "note",
                    self.cell(paragraphs, part),
                    number=k,
                    language=self.language_of(paragraphs),
                )
            )
        return out


def styles_of(fmt) -> frozenset[str]:
    """The styles a run (or a paragraph's own text) shows, from its formatting
    resolved through the styles: bold, italic, underline, struck through,
    superscript, subscript."""
    styles = set()
    if fmt.bold:
        styles.add(STRONG)
    if fmt.italic:
        styles.add(EM)
    if fmt.underline not in (None, "none"):
        styles.add(UNDERLINE)
    if fmt.strike or fmt.double_strike:
        styles.add(STRIKE)
    if fmt.vert_align == "superscript":
        styles.add(SUP)
    elif fmt.vert_align == "subscript":
        styles.add(SUB)
    return frozenset(styles)


def _m(name: str) -> str:
    return f"{{{M_NS}}}{name}"


def _group(text: str) -> str:
    return f"({text})"


def _prop(el, prop: str, name: str, default: str) -> str:
    """The m:val of a property of an equation element, e.g. a bracket."""
    found = el.find(f"{_m(prop)}/{_m(name)}")
    return default if found is None else found.get(_m("val"), default)


def omml_text(el) -> str:
    """An equation (Office Math) as linear text, as pandoc's plain output
    writes it: DV_(it) = α_(i) + β ⋅ X, (a)/(b), √(x), ∑_(i)^(n) x."""
    tag = el.tag.split("}")[-1]

    def part(name: str) -> str:
        found = el.find(_m(name))
        return "" if found is None else omml_text(found)

    if tag == "t":
        return el.text or ""
    if tag.endswith("Pr"):  # properties: not text
        return ""
    if tag == "sSub":
        return part("e") + "_" + _group(part("sub"))
    if tag == "sSup":
        return part("e") + "^" + _group(part("sup"))
    if tag == "sSubSup":
        return part("e") + "_" + _group(part("sub")) + "^" + _group(part("sup"))
    if tag == "f":
        return _group(part("num")) + "/" + _group(part("den"))
    if tag == "rad":
        return "√" + _group(part("e"))
    if tag == "d":
        begin = _prop(el, "dPr", "begChr", "(")
        end = _prop(el, "dPr", "endChr", ")")
        sep = _prop(el, "dPr", "sepChr", "|")
        return begin + sep.join(omml_text(e) for e in el.findall(_m("e"))) + end
    if tag == "nary":
        op = _prop(el, "naryPr", "chr", "∫")
        sub, sup = part("sub"), part("sup")
        return (
            op
            + (f"_{_group(sub)}" if sub else "")
            + (f"^{_group(sup)}" if sup else "")
            + " "
            + part("e")
        )
    text = "".join(omml_text(c) for c in el)
    if tag == "oMath":
        # spaces around relations and operators, which Word does not store
        for op in "=+×⋅<>≤≥≈−":
            text = text.replace(op, f" {op} ")
        text = " ".join(text.split())
    return text


def _notes(document) -> dict:
    """The footnote and endnote elements with the part holding them, by
    (kind, id), separators left out."""
    notes = {}
    for part in document.part.package.iter_parts():
        name = str(part.partname)
        kind = {"/word/footnotes.xml": "footnote", "/word/endnotes.xml": "endnote"}.get(name)
        if kind is None:
            continue
        root = parse_xml(part.blob)
        for note in root.findall(qn(f"w:{kind}")):
            if note.get(qn("w:type")) in (None, "normal"):
                notes[(kind, note.get(qn("w:id")))] = (note, part)
    return notes


def docx_to_markdown(data: bytes, changes: str = "accept") -> str:
    """A Word document's body as Markdown, its tracked changes settled
    ("accept", "reject") or kept as markup ("all"), its comments kept."""
    return to_markdown(read_docx(data, changes))


def read_docx(data: bytes, changes: str = "accept") -> Document:
    """A Word document as prosediff reads it (prosediff.document), its
    tracked changes settled ("accept", "reject") or kept as markup ("all"),
    its comments kept, each paragraph with the language it is marked with."""
    if changes not in CHANGES:
        raise ValueError(f"changes must be one of {CHANGES}, not {changes!r}")
    try:
        document = docx.Document(BytesIO(data))
    except (PackageNotFoundError, KeyError, ValueError) as e:
        raise WordError(str(e) or type(e).__name__) from None
    except Exception as e:  # a zip that is not a Word package, broken XML
        raise WordError(f"{type(e).__name__}: {e}") from None
    reader = Reader(document, changes, languages=WordLanguages(data))
    try:
        reader.comments = {str(c.comment_id): c for c in document.comments}
    except (KeyError, ValueError):
        reader.comments = {}
    reader.notes = _notes(document)
    blocks = reader.body()
    notes = reader.note_blocks()
    return Document(blocks, notes, most_letters(reader.letters))
